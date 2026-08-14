"""
Production Control Board - Phase 1
Reads live from Excel, persists assignments in JSON, ETC timer logic.
Time-study hours and FG completion now sourced from SQL Server.
Worker pool is shift-filtered from live attendance parsing.
"""

from flask import Flask, jsonify, request, render_template
import json, os, sys, queue, threading, time, zipfile, re, webbrowser
from datetime import datetime, date, timedelta
import openpyxl
import warnings
# openpyxl warns on read when it encounters Excel features it doesn't fully
# model (e.g. modern conditional-formatting extensions) — harmless here since
# we only ever read these files, never save them back through openpyxl (Excel
# write-back was disabled early on due to corruption risk). Suppressing just
# openpyxl's own warnings, not warnings globally.
warnings.filterwarnings('ignore', module='openpyxl')

# ── Build stamp ─────────────────────────────────────────────────────────────────
# Bump this whenever app.py changes. It's printed to the console at startup and
# shown in the setup window, so you can confirm at a glance that all three area
# PCs are running the same version — instead of grepping the source or, worse,
# discovering a stale PC only when a fixed bug reappears on one area. Format:
# YYYY.MM.DD[.n] date-based; the trailing note is just a human label.
BUILD_VERSION = '2026.07.30.8'
BUILD_NOTE    = 'printing: strict per-market packing standard (fixes box split), silent Canon path, blocking print-result dialog'

# ── Dev / test mode ─────────────────────────────────────────────────────────────
# When ON, the board:
#   • does NOT write {AREA}_status.json to the shared folder (the big screen never
#     sees your test data), and
#   • treats every SQL WRITE as a dry-run — logged, not executed — so label-print
#     testing can't touch the production [License Plate] table.
# SQL/Excel READS are unaffected, so the board still works normally for testing.
# Default from the BOARD_DEV_MODE env var; the setup window has a checkbox that
# overrides it per launch.
DEV_MODE = os.environ.get('BOARD_DEV_MODE', '').strip().lower() in ('1', 'true', 'yes', 'on')

def sql_write(cursor, sql, params=(), label=''):
    """Run a SQL write, unless DEV_MODE — then log a dry-run and skip it. All
    board-originated INSERT/UPDATE/DELETEs should go through here so dev mode
    reliably keeps test data out of production tables."""
    if DEV_MODE:
        first = sql.strip().splitlines()[0].strip()
        print(f'[dev-mode] SQL write SKIPPED{(" — " + label) if label else ""}: {first} … params={params!r}')
        return None
    return cursor.execute(sql, params)

def _relaunch_without_console():
    """If launched via python.exe (console subsystem), immediately re-exec
    this same script via pythonw.exe (GUI subsystem — no console ever
    attached in the first place) and exit this process.
    This is the reliable fix: ShowWindow-based hiding only hides conhost's
    own window, which on modern Windows Terminal isn't the same thing as
    the visible wt.exe tab — so it can end up looking merely minimized
    instead of actually gone. A real pythonw.exe process sidesteps the
    whole problem. Guarded by an env var so it only ever happens once."""
    if os.name != 'nt':
        return False
    if os.environ.get('_CB_NO_CONSOLE_RELAUNCH'):
        return False  # already relaunched once this chain — don't loop
    exe = sys.executable or ''
    if not exe.lower().endswith('python.exe'):
        return False  # already pythonw.exe or something else — nothing to do
    pythonw = os.path.join(os.path.dirname(exe), 'pythonw.exe')
    if not os.path.exists(pythonw):
        return False  # no pythonw available — caller falls back to ShowWindow
    try:
        import subprocess
        env = os.environ.copy()
        env['_CB_NO_CONSOLE_RELAUNCH'] = '1'
        subprocess.Popen([pythonw, *sys.argv], env=env,
                          creationflags=subprocess.CREATE_NO_WINDOW, close_fds=True)
        os._exit(0)
    except Exception:
        return False

def _hide_console_window():
    """Fallback for when the pythonw.exe relaunch above isn't possible
    (pythonw missing, non-standard install, etc.) — hides (doesn't close)
    this process's own console window via ShowWindow. Less reliable than
    the relaunch on modern Windows Terminal, but better than nothing.
    No-op on non-Windows platforms or if there's no console to hide."""
    if os.name != 'nt':
        return
    try:
        import ctypes
        hwnd = ctypes.windll.kernel32.GetConsoleWindow()
        if hwnd:
            ctypes.windll.user32.ShowWindow(hwnd, 0)  # SW_HIDE
    except Exception:
        pass

# ── Single-instance guard ───────────────────────────────────────────────────────
# Opening the app twice on one PC is genuinely harmful: both processes share the
# same data/ folder, both run the 60s poll AND the 5s dashboard writer, so they
# continuously overwrite each other's assignments/timers JSON (orders "jumping"
# back to the queue, lost worker edits) and fight over port 5050 and the shared
# dashboard status file. This makes sure only one instance of THIS install runs.
_single_instance_handle = None  # kept alive for the process lifetime once acquired

def acquire_single_instance_lock():
    """Returns True if we're the only instance of this install on the PC, False
    if another is already running. Uses a Windows named mutex keyed to this
    install folder — the OS releases it automatically when the owning process
    exits, so there is no stale lock file to clean up even after a crash.
    No-op (returns True) on non-Windows, or if the guard itself can't run — we
    never block startup just because the check is unavailable."""
    global _single_instance_handle
    if os.name != 'nt':
        return True
    try:
        import ctypes, hashlib
        # Keyed to the install path so the SAME folder can't run twice, while a
        # separate install (e.g. a different area) in another folder still can.
        key = hashlib.md5(os.path.abspath(BASE_DIR).lower().encode('utf-8')).hexdigest()[:16]
        name = 'NPV_ProductionControlBoard_' + key  # Local (per-session) namespace — no admin rights needed
        kernel32 = ctypes.windll.kernel32
        handle = kernel32.CreateMutexW(None, False, name)
        ERROR_ALREADY_EXISTS = 183
        if not handle or kernel32.GetLastError() == ERROR_ALREADY_EXISTS:
            if handle:
                kernel32.CloseHandle(handle)
            return False
        _single_instance_handle = handle
        return True
    except Exception as e:
        print(f'[single-instance] guard unavailable ({e}) — continuing without it')
        return True

def _notify_already_running():
    """Tell the user (via a native message box, no Tkinter needed) that the
    board is already open, then the caller exits."""
    print('[single-instance] Another instance of this board is already running on this PC — exiting.')
    if os.name != 'nt':
        return
    try:
        import ctypes
        ctypes.windll.user32.MessageBoxW(
            None,
            'Production Control Board is already running on this PC.\n\n'
            'Only one instance can run at a time — a second one would share the '
            'same data files and make orders jump around on the board.\n\n'
            'Switch to the window that is already open (check the taskbar).',
            'Production Control Board — already running',
            0x40)  # MB_OK | MB_ICONINFORMATION
    except Exception:
        pass

# pyodbc requires the ODBC Driver Manager (unixODBC on Linux, built-in on Windows).
# Import is optional so the rest of the app still runs in environments without it —
# SQL-dependent features (TimeStudy, FG completion) degrade to cached/empty data
# with a clear log message instead of crashing the whole app.
try:
    import pyodbc
    PYODBC_AVAILABLE = True
except ImportError:
    PYODBC_AVAILABLE = False
    print('[SQL] pyodbc not available — TimeStudy and FG completion will use cached/empty data until this is installed and the ODBC driver is present.')

app = Flask(__name__, template_folder='templates', static_folder='static')

# ── Config ─────────────────────────────────────────────────────────────────────
BASE_DIR   = os.path.dirname(os.path.abspath(__file__))
DATA_DIR   = os.path.join(BASE_DIR, 'data')
os.makedirs(DATA_DIR, exist_ok=True)

# File paths — update these to match the actual machine
EXCEL_PLAN = r"\\npvshare\Data\06_Operation\02.Production\02.Component\06. Planning\01. Plan follow up\Packaging_Plan_FollowUp_V01.xlsm"
PLAN_SHEET = '1.FollowPlanForGL'
DATA_START_ROW = 6

# Dev mode: use uploaded files when real paths don't exist
DEV_PLAN = '/mnt/user-data/uploads/Packaging_Plan_FollowUp_V01.xlsm'

AREA  = 'PK'  # placeholder — actual value computed dynamically from EXCEL_PLAN's
              # filename further down (see compute_area_code()), before anything
              # that depends on it (data file names, attendance lookup, stations) runs
AREAS = ['PK', 'RP', 'AD']

# ── SQL Server config ────────────────────────────────────────────────────────
# ASSUMPTION — please verify: TimeStudy lives in TECH_DATA (as given), FG_Database_All
# lives in LOCALNPV (matching the ahead/behind dashboard's server.py). Same server/creds,
# different database. If FG log is actually in a different DB, change SQL_DB_PROD below.
SQL_DRIVER   = '{SQL Server}'  # legacy built-in driver — confirmed via Get-OdbcDriver on this PC (no ODBC Driver 17/13/11 package installed)
SQL_SERVER   = 'VTN1PRDSQL002'
SQL_UID      = 'svcsqllocal'
SQL_PWD      = 'KetnoilocalApp'
SQL_DB_TECH  = 'TECH_DATA'   # TimeStudy table
SQL_DB_PROD  = 'LOCALNPV'    # FG_Database_All table — VERIFY this matches your actual FG log DB

SQL_CONN_TECH = f"DRIVER={SQL_DRIVER};SERVER={SQL_SERVER};DATABASE={SQL_DB_TECH};UID={SQL_UID};PWD={SQL_PWD};"
SQL_CONN_PROD = f"DRIVER={SQL_DRIVER};SERVER={SQL_SERVER};DATABASE={SQL_DB_PROD};UID={SQL_UID};PWD={SQL_PWD};"

FG_POLL_INTERVAL_SEC = 60  # per requirement: refresh FG completion every 1 minute

# Next-shift workers become assignable this many minutes BEFORE their shift
# starts (so GL can pre-staff), but they don't count toward a station's ETC
# until their shift actually begins — see get_available_workers() / 'active'.
NEXT_SHIFT_OPEN_EARLY_HRS = 15 / 60  # 15 minutes

# A PDO is flagged "SAMPLE" when the word "sample" (any case) appears in any of
# these Excel columns. Uses the COL field keys — 'week' is column K (index 10),
# where the sample marker lives. Add more field keys here if needed.
SAMPLE_TEXT_COLS = ('week',)  # column K

# ── Shop floor dashboard feed (shared-drive file, no networking needed) ────────
# Each area board writes its own status snapshot to this shared folder;
# the big-screen dashboard just reads the 3 files — no inbound firewall
# rules needed on any of the 3 production PCs.
DASHBOARD_STATUS_DIR = r"\\npvshare\Data\06_Operation\02.Production\02.Component\06. Planning\dashboard_status"
DASHBOARD_WRITE_INTERVAL_SEC = 5

_dashboard_last_err_ts = 0.0  # throttles repeated dashboard-feed error logging

def write_dashboard_status():
    """Writes this area's status snapshot to the shared dashboard folder as an
    atomic temp-file-then-rename, so the big-screen dashboard never reads a
    half-written file. Hardened for a shared network drive:
      - a UNIQUE temp name per write (PID + ms), so a leftover/locked .tmp from
        a previously-crashed write can never block us (was the Errno 13 case),
      - a short retry on os.replace, which on Windows raises a sharing
        violation (WinError 32) if the dashboard is reading the file at that
        exact instant — it releases within milliseconds, so a couple of retries
        clears it instead of dropping the update,
      - error logging throttled to once a minute so a genuinely unreachable/
        read-only share doesn't flood the console every 5s."""
    global _dashboard_last_err_ts
    if DEV_MODE:
        return  # dev/test mode — never feed the big screen with test data
    try:
        snapshot = build_board_snapshot(AREA)
    except Exception as e:
        print(f'[dashboard feed] snapshot build failed: {e}')
        return
    snapshot['__written'] = int(time.time() * 1000)  # epoch ms — lets a viewer show data age
    body = json.dumps(snapshot, default=str)
    # Two files, same data:
    #   {AREA}_status.json → read server-side by combined_dashboard.py
    #   {AREA}_status.js   → a JSONP wrapper so a plain file:// dashboard.html on the
    #                        share can load it via <script> (fetch() is CORS-blocked
    #                        from file://; <script src> is not), needing no web server
    #                        and no inbound firewall rule anywhere.
    outputs = [
        (f'{AREA}_status.json', body),
        (f'{AREA}_status.js',
         'window.NPV_STATUS=window.NPV_STATUS||{};\n'
         f'window.NPV_STATUS[{AREA!r}]={body};'),
    ]
    try:
        os.makedirs(DASHBOARD_STATUS_DIR, exist_ok=True)
        for fname, text in outputs:
            target_path = os.path.join(DASHBOARD_STATUS_DIR, fname)
            tmp_path = f'{target_path}.{os.getpid()}.{int(time.time()*1000000)}.tmp'
            try:
                with open(tmp_path, 'w', encoding='utf-8') as f:
                    f.write(text)
                last_err = None
                for attempt in range(4):
                    try:
                        os.replace(tmp_path, target_path)  # atomic — no half-written reads
                        last_err = None
                        break
                    except (PermissionError, OSError) as e:
                        last_err = e
                        time.sleep(0.2 * (attempt + 1))  # reader releases within ms
                if last_err is not None:
                    raise last_err
            finally:
                if os.path.exists(tmp_path):
                    try: os.remove(tmp_path)
                    except Exception: pass
    except Exception as e:
        now = time.time()
        if now - _dashboard_last_err_ts > 60:
            print(f'[dashboard feed] write failed (retried; further errors muted 60s): {e}')
            _dashboard_last_err_ts = now

def _dashboard_write_loop():
    while True:
        write_dashboard_status()
        time.sleep(DASHBOARD_WRITE_INTERVAL_SEC)

# ── Attendance / shift config (ported from the ahead/behind dashboard's server.py) ──
CAPACITY_FILES = {
    "AD": r"C:\Users\Nguyen.Nguyen\OneDrive - Northstar Precision Vietnam\Northstar Precision Vietnam - Danh sach nhan su cac khu vuc\Attendance Matrix - Assembly V02.xlsm",
    "PK": r"C:\Users\Nguyen.Nguyen\OneDrive - Northstar Precision Vietnam\Northstar Precision Vietnam - Danh sach nhan su cac khu vuc\Attendance Matrix - Packaging V02.xlsm",
    "RP": r"C:\Users\Nguyen.Nguyen\OneDrive - Northstar Precision Vietnam\Northstar Precision Vietnam - Danh sach nhan su cac khu vuc\Attendance Matrix - Raw Part V02.xlsm",
}

INDIRECT_LABOR_IDS = {
    "NPV18019027", "NPV18021002",   # Assembly GL/TL
    "NPV18018023", "NPV18021011",   # Packaging GL/TL
    "NPV18019028",                  # Raw Part GL/TL
}

# Shift time windows — start/end in decimal hours, per OT level.
# Ported directly from server.py's SHIFTS dict.
SHIFTS = {
    "Ca 1": {
        "start": 6.0, "end": 14.167,
        "ot2_end": 16.167, "ot3_end": 17.167,
    },
    "Ca 2": {
        "start": 14.167, "end": 22.333,
        "ot2_start": 12.167, "ot2_end": 22.333,
        "ot3_start": 11.167, "ot3_end": 22.333,
    },
    "HC": {
        "start": 8.0, "end": 16.5,
        "ot2_end": 18.5, "ot3_end": 19.5,
    },
}

# ── Production hours / shift boundaries (drives ETC freeze + shift-end reset) ────
# Production runs from Ca1 start (06:00) to Ca2 end (22:20); Ca1 and Ca2 overlap
# so there is no gap between them during the day. The only non-production window
# is overnight (22:20 → 06:00), where the ETC timer freezes. Shift boundaries at
# which the ETC is reset+recalculated are Ca1 end (14:10) and the morning start.
PRODUCTION_START = SHIFTS['Ca 1']['start']   # 6.0
CA1_END          = SHIFTS['Ca 1']['end']     # 14.167 (14:10) — Ca1 shift end
PRODUCTION_END   = SHIFTS['Ca 2']['end']     # 22.333 (22:20) — Ca2 shift end / day end

def is_production_time(now=None):
    """True during production hours (06:00–22:20). Outside this the ETC freezes."""
    now = now or datetime.now()
    h = now.hour + now.minute / 60
    return PRODUCTION_START <= h < PRODUCTION_END

def current_shift_phase(now=None):
    """'ca1' (06:00–14:10) | 'ca2' (14:10–22:20) | 'off' (outside production).
    Changing phase is what triggers the shift-end ETC reset."""
    now = now or datetime.now()
    h = now.hour + now.minute / 60
    if h < PRODUCTION_START or h >= PRODUCTION_END:
        return 'off'
    return 'ca1' if h < CA1_END else 'ca2'

ABSENT_CODES = {"NP(Y)", "NP(N)", "NO", "NKLD", "NVR(Y)", "NVR(N)",
                "NP 1/2(Y)", "NP 1/2(N)", "NVR 1/2(Y)", "NVR 1/2(N)", "NĐB"}
HALF_DAY_CODES = {"NP 1/2(Y)", "NP 1/2(N)", "NVR 1/2(Y)", "NVR 1/2(N)"}

DEV_ATTENDANCE = '/mnt/user-data/uploads/attendance_sample.xlsm'  # not provided yet — see notes

# ── Column indices (0-based) from row 5 headers ────────────────────────────────
COL = {
    'pdo':          0,   # A  PO
    'part':         1,   # B  Part Number
    'raw_part':     2,   # C  Raw Part
    'desc':         3,   # D  PartDescription
    'qty':          4,   # E  Order Q'ty
    'dest':         5,   # F  Destination
    'pack_type':    6,   # G  NPV Pack/Assemble
    'receive_date': 7,   # H  Ngày VPIC1 có sẵn hàng (date goods became available)
    'ship_date':    8,   # I  Ship date
    'ship_mode':    9,   # J  Ship mode
    'week':         10,  # K  Week
    'note':         11,  # L  Note
    'docs':         12,  # M  Docs
    'vendor_name':  13,  # N  Nhà cung cấp (full vendor/supplier name)
    'pii_po':       16,  # Q  PII PO
    'received':     14,  # O  Received
    'progress':     15,  # P  Progress
    'finished_hrs': 17,  # R  Finished hrs
    'vendor':       18,  # S  Vendor code
    'station_t':    19,  # T  Station (thực tế) — existing assignment
    'boxes':        20,  # U  Số thùng nhỏ — free-text carton breakdown note, not a count
    'station_w':    22,  # W  Station (web board writes here)
}

# ── Filename-based area safeguard ───────────────────────────────────────────────
# The plan Excel can contain rows from multiple areas (column G / "NPV Pack/Assemble").
# The file name itself tells us which area this board is for, and we use that to
# filter column G — this stops the board from silently showing another area's
# orders if the shared file has stray cross-area rows, AND acts as a hard safety
# check: if the file name doesn't match a known convention, we genuinely don't
# know which area's data to trust, so the server must refuse to start rather than
# guess.
FILE_AREA_PREFIXES = {
    'Packaging': 'Packaging',
    'AssyD':     'Assembly',
    'RawPart':   'Raw Part',
}

def detect_area_filter(excel_path):
    """Returns (area_filter, error_message).
    area_filter is the substring to match against column G — None if the file
    name doesn't start with a recognized prefix, in which case error_message
    explains why and the server must not start."""
    if not excel_path:
        return None, 'No plan file path configured.'
    fname = os.path.basename(excel_path)
    for prefix, area in FILE_AREA_PREFIXES.items():
        if fname.startswith(prefix):
            return area, None
    return None, (
        f'File name "{fname}" does not start with a recognized prefix '
        f'(Packaging / AssyD / RawPart) — cannot determine which area\'s '
        f'orders belong on this board.'
    )

AREA_DISPLAY_NAMES = {
    'Packaging': 'PACKAGING',
    'Assembly':  'ASSEMBLY',
    'Raw Part':  'RAW PART',
}

# 2-letter area codes — matches the station-naming convention used in SQL
# (AD01, PK02, RP03) and everywhere else in this app (data file names,
# attendance file lookup, the Flask route). This is THE key that makes one
# codebase correctly serve any of the 3 areas depending only on which Excel
# file it's pointed at — get it wrong and stations/data files/attendance all
# silently point at the wrong area.
AREA_CODE_MAP = {
    'Packaging': 'PK',
    'Assembly':  'AD',
    'Raw Part':  'RP',
}

def compute_area_code():
    """Determines this instance's 2-letter area code from the currently
    configured EXCEL_PLAN file name. Falls back to 'PK' if undetermined —
    harmless, since the invalid-filename hard-block elsewhere refuses to
    start the server in that case anyway, so this fallback is never actually
    reached in a running instance."""
    area_filter, _ = detect_area_filter(EXCEL_PLAN)
    if area_filter is None:
        return 'PK'
    return AREA_CODE_MAP.get(area_filter, 'PK')

def normalize_station_sql(station_id):
    """Board format 'PK-01' -> SQL format 'PK01' (no dash)."""
    return (station_id or '').replace('-', '').upper()

def normalize_station_board(sql_station):
    """SQL format 'PK01' -> board format 'PK-01' (insert dash after 2-letter prefix).
    Tolerates a trailing suffix that a second label printer adds for the small
    inner boxes, e.g. 'RP10-ThungNho' -> 'RP-10' — only the leading
    <2 letters><digits> station code is used."""
    s = (sql_station or '').strip().upper()
    m = re.match(r'^([A-Z]{2})(\d+)', s)
    if m:
        return m.group(1) + '-' + m.group(2)
    return s


# ── In-memory cache ────────────────────────────────────────────────────────────
_raw_pdo_cache   = []   # from Excel only: qty, dest, ship_date, received, part, + time_study fields
_pdo_cache       = []   # FINAL merged (raw + SQL FG completion), filtered, sorted — what the API serves
_worker_cache    = []   # today's full on-shift roster; time-of-day filtering (pre-open/active/shift-end) is applied live per board GET
_shift_windows_cache = {
    'mode': 'split', 'ca1_start': 6.0, 'ca1_end': 14.167,
    'ca2_start': 14.167, 'ca2_end': 22.333, 'source': 'fallback_static',
}
_last_worker_refresh_ts = 0
_day_capacity_hours = 0.0  # total available labour-hours today (from attendance) — for the report chart line
WORKER_REFRESH_THROTTLE_SEC = 300  # attendance file read is relatively heavy — check every 5 minutes
_station_cache   = []
_time_study_cache = {}  # {part_number: {'time_study':.., 'ops_qty':..}}
_fg_completed_cache = {}  # {po: completed_qty} from SQL, refreshed every 60s
_split_completed_cache = {}  # {(parent_id, letter): live_completed_qty} from SQL, refreshed every 60s
_wip_cache = {}  # {(po, sql_station): (wip_qty, timestamp)} from Nhaplecuoingay_All, refreshed every 60s
_fg_latest_scan_cache = {}  # {(po, sql_station): latest_input_time} from FG_Database_All, refreshed every 60s
_license_plate_cache = {}  # {po: (sql_station, print_time)} from License Plate, refreshed every 60s
_cache_lock      = threading.RLock()
_last_excel_mtime = 0
_last_shift_phase = None   # 'ca1'|'ca2'|'off' — last seen, to detect shift-end boundaries
_last_pdo_status  = 'ok'   # 'ok' | 'locked' | 'corrupted' | 'error' | 'invalid_filename'
# True only after TODAY's attendance file was read cleanly. An empty roster from a
# failed/missing/unsynced read ([] on error) is indistinguishable from a genuine
# "nobody on shift", so this flag gates prune_expired_workers: never strip assigned
# operators off the board when we simply couldn't read attendance (the mass-wipe bug).
_last_attendance_ok = True
_area_label       = 'PRODUCTION'  # display header — set from detected filename area
_last_fg_status   = 'ok'   # 'ok' | 'no_pyodbc' | 'error'
_last_ts_status   = 'ok'

# ── Persistence helpers ────────────────────────────────────────────────────────
def _path(name):
    return os.path.join(DATA_DIR, name)

def load_json(name, default):
    p = _path(name)
    if not os.path.exists(p):
        return default
    try:
        with open(p, encoding='utf-8') as f:
            return json.load(f)
    except (json.JSONDecodeError, UnicodeDecodeError, OSError) as e:
        # A truncated/corrupt or wrong-encoding state file must not take the
        # whole board (or the 60s poll loop) down — fall back to the default
        # and log it. save_json always writes UTF-8, so this also recovers a
        # file written by an older build under a non-UTF-8 Windows locale.
        print(f'[load_json error] {name}: {e} — using default')
        return default

def save_json(name, data):
    with open(_path(name), 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def get_fallback_stations():
    """Default station IDs for THIS instance's area — must be computed live
    (not a fixed list) since AREA is only known after detect_area_filter()
    runs against the configured Excel file."""
    return [f'{AREA}-{str(i).zfill(2)}' for i in range(1, 13)]

def assignments_path(area):
    return _path(f'assignments_{area}.json')

def load_assignments():
    path = assignments_path(AREA)
    stations = _station_cache if _station_cache else get_fallback_stations()
    if not os.path.exists(path):
        data = {s: {'inprogress': None, 'queue': []} for s in stations}
        save_assignments(data)
        return data
    try:
        with open(path, encoding='utf-8') as f:
            data = json.load(f)
    except (json.JSONDecodeError, UnicodeDecodeError, OSError) as e:
        # Corrupt/unreadable assignments file — reseed empty rather than crash
        # every caller (poll loop, board API). Better a cleared board than a
        # server that won't serve at all.
        print(f'[load_assignments error] {e} — reseeding empty assignments')
        data = {s: {'inprogress': None, 'queue': []} for s in stations}
        save_assignments(data)
        return data
    changed = False
    for s in stations:
        if s not in data:
            data[s] = {'inprogress': None, 'queue': []}
            changed = True
        # Migrate old {inprogress, next} format to {inprogress, queue}
        if 'queue' not in data[s]:
            old_next = data[s].pop('next', None)
            data[s]['queue'] = [old_next] if old_next else []
            changed = True
    if changed:
        save_assignments(data)
    return data

def save_assignments(data):
    save_json(f'assignments_{AREA}.json', data)

def load_workers_assigned():
    """{ station_id: [worker_id, ...] }"""
    return load_json(f'workers_{AREA}.json', {})

def save_workers_assigned(data):
    save_json(f'workers_{AREA}.json', data)

def load_timers():
    """{ station_id: { started_at, remaining_hrs, total_hrs, worker_count } }"""
    return load_json(f'timers_{AREA}.json', {})

def save_timers(data):
    save_json(f'timers_{AREA}.json', data)

def load_splits():
    """{ parent_pdo_id: { splits: { 'A': {station, station_set_at, credit, total_qty}, 'B': {...} } } }"""
    return load_json(f'splits_{AREA}.json', {})

def save_splits(data):
    save_json(f'splits_{AREA}.json', data)

def parse_split_id(pdo_id):
    """Returns (parent_id, letter) if pdo_id is a real, currently-active split
    child (verified against splits_{AREA}.json, not just pattern-matched — PDO IDs
    already contain hyphens, e.g. 'PDO-421010', so a naive regex match alone
    isn't enough to safely tell a split apart from a normal PDO)."""
    m = re.match(r'^(.+)-([ABCD])$', pdo_id)
    if not m:
        return None, None
    parent_id, letter = m.group(1), m.group(2)
    splits_data = load_splits()
    if parent_id in splits_data and letter in splits_data[parent_id].get('splits', {}):
        return parent_id, letter
    return None, None

def update_split_station(pdo_id, new_station):
    """Called whenever a PDO gets assigned/moved to a station. If pdo_id is a
    split child, update its tracked station and reset its live-tracking
    cutoff to now — this is what lets a split that started unassigned (sitting
    in the pool) begin accurate SQL attribution the moment GL actually places
    it, without retroactively picking up unrelated scans from before it
    arrived there."""
    parent_id, letter = parse_split_id(pdo_id)
    if not parent_id:
        return
    splits_data = load_splits()
    rec = splits_data[parent_id]['splits'][letter]
    if rec.get('station') != new_station:
        rec['station'] = new_station
        rec['station_set_at'] = datetime.now().isoformat()
        save_splits(splits_data)

# ── Excel readers ──────────────────────────────────────────────────────────────
def _resolve_path(real, dev):
    return real if os.path.exists(real) else dev

def _is_excel_locked(path):
    """Excel creates a ~$filename lock file whenever the file is open —
    used only for status messaging now, not as a hard block on reading."""
    folder = os.path.dirname(path)
    fname  = os.path.basename(path)
    lock   = os.path.join(folder, '~$' + fname)
    return os.path.exists(lock)

def _safe_load_workbook(path, **kwargs):
    """Reads the Excel file via a fresh temp copy, retrying briefly on failure.

    Excel only holds an exclusive write-lock for the brief moment it's actually
    saving (Ctrl+S, autosave, or on close) — not continuously while the file is
    open and being viewed/edited. A read that happens to land exactly during
    that save can get a 'torn' read (incomplete/corrupt zip data). Rather than
    refusing to read at all whenever Excel has the file open, we copy the file
    first (fast, reduces the race window) and retry a few times if it fails —
    since those save moments are brief and self-resolve within ~1 second.

    Raises the last exception if all retries are exhausted.
    """
    import shutil, tempfile, time as _time
    last_err = None
    for attempt in range(3):
        tmp_fd, tmp_path = tempfile.mkstemp(suffix=os.path.splitext(path)[1])
        os.close(tmp_fd)
        try:
            shutil.copy2(path, tmp_path)
            wb = openpyxl.load_workbook(tmp_path, **kwargs)
            return wb, tmp_path  # caller must os.remove(tmp_path) after use
        except (zipfile.BadZipFile, OSError, PermissionError) as e:
            last_err = e
            try: os.remove(tmp_path)
            except: pass
            if attempt < 2:
                _time.sleep(0.4 * (attempt + 1))  # 0.4s, 0.8s backoff
                continue
        except Exception as e:
            last_err = e
            try: os.remove(tmp_path)
            except: pass
            break
    raise last_err

def _parse_box_note(val):
    """Column U ('Số thùng nhỏ') is a free-text carton/box breakdown note —
    e.g. '0Tx120 pcs; 1Tx5 pcs' — not a number. Just return it as trimmed
    text; there is nothing here to sum."""
    if val is None:
        return ''
    return str(val).strip()

def _row_is_sample(row):
    """True if the word "sample" (any case) appears in any of the configured
    Excel columns (SAMPLE_TEXT_COLS) for this row — used to flag sample orders."""
    for field in SAMPLE_TEXT_COLS:
        idx = COL.get(field)
        if idx is not None and idx < len(row):
            v = row[idx]
            if v and 'sample' in str(v).lower():
                return True
    return False

def read_pdos():
    """Reads RAW PDO rows from Excel: qty, dest, ship_date, received, part number.
    Does NOT filter by completion — that now depends on live SQL FG data and
    happens in rebuild_final_pdo_cache(). Only the 'received > 0' (material in)
    filter stays here since that's still Excel-sourced.
    Also enforces the filename→area safeguard: rows whose column G doesn't match
    the area implied by the file name are excluded, and if the file name itself
    doesn't match a known convention, refuses to return any data at all.
    Returns (pdos, status): 'ok', 'locked', 'corrupted', 'error', 'invalid_filename'."""
    path = _resolve_path(EXCEL_PLAN, DEV_PLAN)

    area_filter, area_err = detect_area_filter(path)
    if area_filter is None:
        print(f'[INVALID FILENAME] {area_err}')
        print('[INVALID FILENAME] Server will not serve PDO data until the plan file path is fixed.')
        return [], 'invalid_filename'

    global _area_label
    _area_label = AREA_DISPLAY_NAMES.get(area_filter, area_filter.upper())

    tmp_path = None
    try:
        wb, tmp_path = _safe_load_workbook(path, read_only=True, data_only=True)
        ws = wb[PLAN_SHEET]
        pdos = []
        today = date.today()
        for row in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
            pdo_id = row[COL['pdo']]
            if not pdo_id:
                break
            pdo_id = str(pdo_id)
            if not pdo_id.startswith('PDO'):
                continue
            pack_type_val = str(row[COL['pack_type']] or '')
            if area_filter.lower() not in pack_type_val.lower():
                continue  # belongs to a different area — skip

            qty = row[COL['qty']] or 0
            if qty <= 0:
                continue

            ship_raw = row[COL['ship_date']]
            if not isinstance(ship_raw, datetime):
                # Ship date column has no real date value (blank, text like
                # "TBD", etc.) — can't compute urgency/priority for this order
                # at all, so it's excluded rather than guessed at with a fake
                # fallback days_left.
                continue
            ship_str  = ship_raw.strftime('%Y-%m-%d')
            days_left = (ship_raw.date() - today).days

            received = row[COL['received']] or 0
            # No date-window restriction: ALL orders in the plan are loaded
            # (regardless of ship date or whether material is in yet). Finished
            # orders are still dropped later from the active list by
            # rebuild_final_pdo_cache (SQL FG completion) and only surface in the
            # separate Finished report; the queue/PDO-report selection therefore
            # shows every unfinished order in the plan.

            recv_raw = row[COL['receive_date']]
            recv_str = recv_raw.strftime('%Y-%m-%d') if isinstance(recv_raw, datetime) else ''

            boxes = _parse_box_note(row[COL['boxes']])

            part = str(row[COL['part']] or '')
            pdos.append({
                'id':           pdo_id,
                'part':         part,
                'raw_part':     str(row[COL['raw_part']] or ''),
                'desc':         str(row[COL['desc']] or ''),
                'qty':          qty,
                'dest':         str(row[COL['dest']] or ''),
                'pack_type':    str(row[COL['pack_type']] or ''),
                'ship_date':    ship_str,
                'receive_date': recv_str,
                'days_left':    days_left,
                'ship_mode':    str(row[COL['ship_mode']] or ''),
                'week':         str(row[COL['week']] or ''),
                'note':         str(row[COL['note']] or '') if row[COL['note']] else '',
                'docs':         str(row[COL['docs']] or ''),
                'pii_po':       str(row[COL['pii_po']] or ''),
                'received':     received,
                'vendor':       str(row[COL['vendor']] or ''),
                'vendor_name':  str(row[COL['vendor_name']] or ''),
                'boxes':        boxes,
                'is_sample':    _row_is_sample(row),
            })
        wb.close()

        # ── Attach TimeStudy fields (SQL) ──────────────────────────────────
        unique_parts = list({p['part'] for p in pdos if p['part']})
        ts_map = get_time_study(unique_parts)
        for p in pdos:
            ts = ts_map.get(p['part'])
            if ts and ts['time_study'] > 0:
                p['time_study'] = ts['time_study']
                p['ops_qty']    = ts['ops_qty']  # kept for reference only; no longer used in the hours math
                # Hours are now always expressed for a SINGLE operator to do the
                # whole order (qty / time_study). This is both the ETC baseline
                # (ETC then scales down by the actual operators staffed on the
                # station) AND the "required hours to complete" shown to GL —
                # the standard-crew (ops_qty) divisor was removed per request.
                p['finished_hrs'] = round(p['qty'] / ts['time_study'], 2)
                p['std_hours']    = p['finished_hrs']
            else:
                p['time_study']   = 0
                p['ops_qty']      = 1
                p['finished_hrs'] = 0
                p['std_hours']    = 0
            # Urgency uses the same 1-operator required-hours figure.
            p['urgency'] = round(p['days_left'] * 10 - p['std_hours'], 1)

        return pdos, 'ok'
    except zipfile.BadZipFile as e:
        # If Excel currently has the file open, this was very likely a torn
        # read during a save — transient, not real corruption.
        if _is_excel_locked(path):
            print(f'[Excel busy] Read failed while Excel had the file open (likely mid-save): {e}. Using cached data — will retry on next refresh.')
            return _raw_pdo_cache, 'locked'
        print(f'[CORRUPTED FILE] {path}')
        print(f'[CORRUPTED FILE] openpyxl cannot read this .xlsm even with Excel closed — the file itself may be damaged: {e}')
        print(f'[CORRUPTED FILE] Try opening it directly in Excel. If Excel offers to "repair", the file is corrupted.')
        print(f'[CORRUPTED FILE] Restore from OneDrive/Desktop version history if needed.')
        return _raw_pdo_cache, 'corrupted'
    except Exception as e:
        print(f'[Excel error] {e}')
        return _raw_pdo_cache, 'error'
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try: os.remove(tmp_path)
            except: pass

# ── SQL: TimeStudy lookup ───────────────────────────────────────────────────────
def get_time_study(part_numbers):
    """Returns {part_number: {'time_study': pcs/hr for 1 worker, 'ops_qty': standard crew size}}.
    ASSUMPTION: TimeStudy.FG matches the Excel Part Number column as an exact string —
    verify this join key if results come back empty for parts you know exist."""
    global _last_ts_status
    if not part_numbers:
        return {}
    if not PYODBC_AVAILABLE:
        _last_ts_status = 'no_pyodbc'
        return _time_study_cache
    try:
        conn = pyodbc.connect(SQL_CONN_TECH, timeout=10)
        cursor = conn.cursor()
        placeholders = ','.join('?' for _ in part_numbers)
        cursor.execute(
            f"SELECT FG, Time_Study, Ops_Qty FROM [dbo].[TimeStudy] WHERE FG IN ({placeholders})",
            *part_numbers
        )
        rows = cursor.fetchall()
        conn.close()
        result = {}
        for r in rows:
            fg = str(r[0]).strip()
            ts = float(r[1]) if r[1] else 0
            ops = int(r[2]) if r[2] else 1
            result[fg] = {'time_study': ts, 'ops_qty': max(ops, 1)}
        _time_study_cache.update(result)
        _last_ts_status = 'ok'
        return _time_study_cache
    except Exception as e:
        print(f'[TimeStudy SQL error] {e}')
        _last_ts_status = 'error'
        return _time_study_cache

# ── SQL: FG completion polling ──────────────────────────────────────────────────
def get_fg_completed(po_numbers=None):
    """Returns {po: total_completed_qty} — sum of Finished Qty per PO, all-time.
    ASSUMPTION: PO numbers are never reused across different orders, so an
    all-time SUM is safe. If POs can recur, this needs a date-range filter.

    When `po_numbers` is given, the SUM is restricted to just those POs
    (WHERE PO IN (...)) instead of grouping the ENTIRE FG_Database_All log
    every time. The board only ever shows a few hundred PDOs, so this turns a
    full-table all-time GROUP BY (the single biggest cost in a refresh on a
    large FG log) into a small keyed lookup. Same parameterized-IN pattern
    get_time_study() already uses. Falls back to the full scan only if no PO
    list is provided (kept for safety/back-compat)."""
    global _last_fg_status
    if not PYODBC_AVAILABLE:
        _last_fg_status = 'no_pyodbc'
        return _fg_completed_cache
    try:
        conn = pyodbc.connect(SQL_CONN_PROD, timeout=10)
        cursor = conn.cursor()
        result = {}
        pos = [str(p).strip() for p in (po_numbers or []) if p]
        if po_numbers is not None and not pos:
            # Board is empty — nothing to look up, skip the query entirely.
            conn.close()
            _last_fg_status = 'ok'
            return result
        if pos:
            # Chunk to stay under SQL Server's ~2100-parameter cap on large boards.
            for i in range(0, len(pos), 1000):
                chunk = pos[i:i+1000]
                placeholders = ','.join('?' for _ in chunk)
                cursor.execute(
                    f"SELECT PO, SUM([Finished Qty]) AS completed "
                    f"FROM [dbo].[FG_Database_All] WHERE PO IN ({placeholders}) GROUP BY PO",
                    *chunk)
                for r in cursor.fetchall():
                    result[str(r[0]).strip()] = int(r[1]) if r[1] else 0
        else:
            cursor.execute("""
                SELECT PO, SUM([Finished Qty]) AS completed
                FROM [dbo].[FG_Database_All]
                GROUP BY PO
            """)
            result = {str(r[0]).strip(): int(r[1]) if r[1] else 0 for r in cursor.fetchall()}
        conn.close()
        _last_fg_status = 'ok'
        return result
    except Exception as e:
        print(f'[FG SQL error] {e}')
        _last_fg_status = 'error'
        return _fg_completed_cache

def get_fg_latest_scan_times():
    """Returns {(po, sql_station): latest_input_time} — the most recent
    FULL-BOX scan timestamp per PO+station in FG_Database_All. Used to tell
    whether a Nhaplecuoingay_All (end-of-shift WIP) entry has since been
    absorbed by a subsequent full-box scan at that same station — if so, the
    WIP number is stale and must NOT also be subtracted, or completed qty
    would be double-counted. Windowed to the last 7 days to keep this cheap;
    WIP deduction only ever matters for whatever's currently in-progress."""
    if not PYODBC_AVAILABLE:
        return {}
    try:
        conn = pyodbc.connect(SQL_CONN_PROD, timeout=10)
        cursor = conn.cursor()
        cursor.execute("""
            SELECT PO, Station, MAX([Input time]) AS last_scan
            FROM [dbo].[FG_Database_All]
            WHERE [Input time] >= ?
            GROUP BY PO, Station
        """, datetime.now() - timedelta(days=7))
        rows = cursor.fetchall()
        conn.close()
        return {(str(r[0]).strip(), str(r[1]).strip().upper()): r[2] for r in rows if r[2]}
    except Exception as e:
        print(f'[FG latest-scan SQL error] {e}')
        return None  # None (not {}) signals "query failed" so the caller keeps the last good cache

def get_wip_entries():
    """Returns {(po, sql_station): (wip_qty, timestamp)} — the most recent
    end-of-shift WIP entry per PO+station from Nhaplecuoingay_All (operators
    log partial/incomplete box qty here at shift end). Same column shape as
    FG_Database_All (PO, Station, [Finished Qty], [Input time]).
    Rows come back ordered oldest-first, so each dict write naturally keeps
    only the LATEST entry per PO+station — no separate max() step needed."""
    if not PYODBC_AVAILABLE:
        return {}
    try:
        conn = pyodbc.connect(SQL_CONN_PROD, timeout=10)
        cursor = conn.cursor()
        cursor.execute("""
            SELECT PO, Station, [Finished Qty], [Input time]
            FROM [dbo].[Nhaplecuoingay_All]
            WHERE [Input time] >= ?
            ORDER BY [Input time] ASC
        """, datetime.now() - timedelta(days=7))
        rows = cursor.fetchall()
        conn.close()
        result = {}
        for po, station, qty, ts in rows:
            if not po or not station:
                continue
            result[(str(po).strip(), str(station).strip().upper())] = (int(qty) if qty else 0, ts)
        return result
    except Exception as e:
        print(f'[Nhaplecuoingay SQL error] {e}')
        return None  # None (not {}) signals "query failed" so the caller keeps the last good cache

def get_wip_deduction(po_id, station_id):
    """Returns the extra qty to count as 'completed' for this PO at this
    station, based on the most recent Nhaplecuoingay_All (end-of-shift WIP)
    entry there — but ONLY if no full-box FG_Database_All scan has happened
    at that same station since that WIP entry was logged. A later full-box
    scan means the WIP was already absorbed into the real completed count,
    so counting it again would double-subtract from remaining_qty."""
    if not station_id:
        return 0
    key = (str(po_id).strip(), normalize_station_sql(station_id))
    wip = _wip_cache.get(key)
    if not wip:
        return 0
    wip_qty, wip_ts = wip
    last_scan_ts = _fg_latest_scan_cache.get(key)
    if last_scan_ts and last_scan_ts > wip_ts:
        return 0  # absorbed by a subsequent full-box scan — don't double-count
    return wip_qty

def get_wip_deduction_latest(po_id):
    """Partial-box WIP progress for a PO, taken from the MOST RECENT shift-end
    WIP entry across whatever station last logged it — not just the station the
    order currently sits on.

    Why latest-across-stations rather than the current station only: a WIP entry
    is real completed work on the PO. A single order is worked at one station at
    a time; as it moves (X → Y → back to X) each station's shift-end log is just
    a newer snapshot of the SAME partial units, so the newest one is the truth.
    Keying the deduction to the order's *current* station (the old behaviour)
    made that progress vanish the moment the order was reassigned and not
    reliably come back — this fixes that. The per-station 'already absorbed by a
    later full-box FG scan' guard still applies to each candidate entry."""
    po = str(po_id).strip()
    best_qty, best_ts = 0, None
    for (wpo, wstation), (wip_qty, wip_ts) in _wip_cache.items():
        if wpo != po:
            continue
        last_scan_ts = _fg_latest_scan_cache.get((wpo, wstation))
        if last_scan_ts and wip_ts is not None and last_scan_ts > wip_ts:
            continue  # this station's WIP already folded into the full-box FG count
        if best_ts is None or (wip_ts is not None and wip_ts > best_ts):
            best_qty, best_ts = wip_qty, wip_ts
    return best_qty

def get_license_plate_prints():
    """Returns {po: (sql_station, print_time)} — the SINGLE most recent print
    record per PO within the last 30 days. A GL reprinting a label can change
    which station it's destined for, so only the latest print counts, not
    every print event. Same last-write-wins dedup pattern as get_wip_entries()."""
    if not PYODBC_AVAILABLE:
        return {}
    try:
        conn = pyodbc.connect(SQL_CONN_PROD, timeout=10)
        cursor = conn.cursor()
        cursor.execute("""
            SELECT [PO number], [Work Station], [Print time]
            FROM [dbo].[License Plate]
            WHERE [Print time] >= ?
            ORDER BY [Print time] ASC
        """, datetime.now() - timedelta(days=30))
        rows = cursor.fetchall()
        conn.close()
        result = {}
        latest_ts = None
        for po, station, ts in rows:
            if not po or not station:
                continue
            result[str(po).strip()] = (str(station).strip().upper(), ts)
            if ts is not None and (latest_ts is None or ts > latest_ts):
                latest_ts = ts
        # One-line health signal, logged only when it CHANGES (not every 60s poll):
        # shows the query is alive, how many distinct POs it sees, and the most
        # recent print time. If a freshly-printed order isn't auto-placing, this
        # tells you at a glance whether the LP query is even seeing it — if
        # latest_print stops advancing while labels are still printing, the app
        # is querying the wrong table/columns or a stale mirror.
        summary = (len(result), str(latest_ts))
        if getattr(get_license_plate_prints, '_last_summary', None) != summary:
            get_license_plate_prints._last_summary = summary
            print(f'[License Plate] query ok — {len(result)} PO(s) in the 30-day window, '
                  f'latest print {latest_ts}')
        return result
    except Exception as e:
        print(f'[License Plate SQL error] {e}')
        return None  # None (not {}) signals "query failed" so the caller keeps the last good cache

def _record_lp_suppression(po):
    """After a MANUAL removal/unassign of an auto-queued PDO, remember the
    License Plate print timestamp we are suppressing up to. auto_queue_from_
    license_plate() then won't re-add that PDO until a NEWER LP print is issued
    for it (compared by timestamp)."""
    if not po:
        return
    rec = _license_plate_cache.get(str(po).strip())
    if not rec:
        return  # not an LP-driven order — nothing to suppress
    _station, print_time = rec
    try:
        supp = load_json(f'lp_suppress_{AREA}.json', {})
        supp[str(po).strip()] = print_time.isoformat() if hasattr(print_time, 'isoformat') else str(print_time)
        save_json(f'lp_suppress_{AREA}.json', supp)
    except Exception as e:
        print(f'[lp-suppress error] {e}')

def get_split_live_completed(parent_pdo_id, station_id, cutoff_iso):
    """SQL sum for ONE split: same PO, but only scans at the split's CURRENT
    station, and only scans that happened after that station was set (the
    cutoff moves forward every time the split gets placed/replaced on a
    station — not just at creation). That's what stops double-counting: if a
    split lands on a station that already had unrelated scans (either from
    before the split existed, or from other work), only scans from the moment
    THIS split arrived there count."""
    if not station_id or not cutoff_iso:
        return 0
    if not PYODBC_AVAILABLE:
        return 0
    try:
        conn = pyodbc.connect(SQL_CONN_PROD, timeout=10)
        cursor = conn.cursor()
        sql_station = normalize_station_sql(station_id)
        # Bind the cutoff as a real datetime, not the raw ISO string. The stored
        # station_set_at is datetime.now().isoformat() (T-separator + 6 microsecond
        # digits), which SQL Server can't implicitly convert to datetime for the
        # [Input time] comparison — it throws "Conversion failed …" (error 241) and
        # the whole per-split completion silently reads 0. pyodbc binds a datetime
        # object cleanly (same as the License Plate / FG queries).
        cutoff = cutoff_iso
        if isinstance(cutoff, str):
            try:
                cutoff = datetime.fromisoformat(cutoff)
            except Exception:
                cutoff = cutoff_iso  # last resort; will still error but stays logged
        cursor.execute("""
            SELECT SUM([Finished Qty]) AS completed
            FROM [dbo].[FG_Database_All]
            WHERE PO = ? AND Station = ? AND [Input time] > ?
        """, parent_pdo_id, sql_station, cutoff)
        row = cursor.fetchone()
        conn.close()
        return int(row[0]) if row and row[0] else 0
    except Exception as e:
        print(f'[split SQL error] {parent_pdo_id}/{station_id}: {e}')
        return 0

def refresh_split_completions():
    """Query live completion for every active split — each using its OWN
    current station + station-set-timestamp cutoff (both may be None if the
    split hasn't been placed on a station yet, sitting unassigned in the
    pool — in that case there's nothing to query, completion is credit-only).
    Runs on the same 60s cycle as the regular FG poll."""
    global _split_completed_cache
    splits_data = load_splits()
    new_cache = {}
    for parent_id, info in splits_data.items():
        for letter, s in info.get('splits', {}).items():
            live = get_split_live_completed(parent_id, s.get('station'), s.get('station_set_at'))
            new_cache[(parent_id, letter)] = live
    _split_completed_cache = new_cache

def inject_splits_into_cache(merged, assignments):
    """Adds synthetic split-PDO entries (PDO-xxx-A, -B, ...) to the board's
    active list, and reduces (or removes) the original parent entry by
    whatever's been carved off into active splits — the 'leftover' concept.
    Each split's own remaining_qty comes from its one-time credit (inherited
    only if it shares the PDO's original station) plus its own live SQL
    tracking, station+timestamp attributed so reused stations don't
    double-count pre-split scans. WIP (Nhaplecuoingay_All) deduction is
    applied per-split using its own current station, and to the leftover
    parent using whatever station it's actually assigned to on the board —
    needs `assignments` since the leftover recompute below fully overwrites
    whatever rebuild_final_pdo_cache() already set for that entry."""
    splits_data = load_splits()
    if not splits_data:
        return merged

    by_id = {p['id']: p for p in merged}
    raw_by_id = {p['id']: p for p in _raw_pdo_cache}

    for parent_id, info in splits_data.items():
        parent_static = raw_by_id.get(parent_id) or by_id.get(parent_id)
        if not parent_static:
            continue  # parent metadata unavailable — can't render splits for it

        parent_qty = parent_static['qty']
        parent_completed_global = _fg_completed_cache.get(parent_id, 0)
        parent_completed_global = max(0, min(parent_completed_global, parent_qty))
        global_remaining = parent_qty - parent_completed_global

        # 1st pass: each split's OWN tracked completion (credit + station-attributed
        # live SQL + per-split WIP).
        rows = []  # [letter, s, completed, remaining]
        for letter, s in info.get('splits', {}).items():
            live = _split_completed_cache.get((parent_id, letter), 0)
            completed = min(s['credit'] + live, s['total_qty'])
            remaining = max(0, s['total_qty'] - completed)
            wip_ded = get_wip_deduction(parent_id, s.get('station'))
            if wip_ded:
                extra = min(wip_ded, remaining)
                completed += extra
                remaining -= extra
            rows.append([letter, s, completed, remaining])

        # Reconcile against the AUTHORITATIVE PO-level FG completion. Station+cutoff
        # attribution can miss real completion — scans logged under a different
        # station string (e.g. a printer suffix), or before a split's cutoff because
        # it was re-placed after the work. That completion still belongs to this PO,
        # so credit whatever the per-split tracking didn't account for to the splits
        # in order. Without this, a PO that is fully complete in SQL leaves its
        # splits stuck on stations at 0% — they never reach remaining 0 to flush,
        # while the parent (correctly) drops into the Finished report.
        attributed = sum(r[2] for r in rows)
        unattributed = max(0, parent_completed_global - attributed)
        for r in rows:
            if unattributed <= 0:
                break
            take = min(unattributed, r[3])
            r[2] += take; r[3] -= take; unattributed -= take

        active_remaining_sum = 0
        for letter, s, completed, remaining in rows:
            active_remaining_sum += remaining
            if remaining > 0:
                split_id = f"{parent_id}-{letter}"
                sp = dict(parent_static)
                sp['id'] = split_id
                sp['is_split_child'] = True
                sp['split_letter'] = letter
                sp['split_station'] = s.get('station')  # may be None — unassigned, sitting in pool
                sp['parent_id'] = parent_id
                sp['qty'] = s['total_qty']
                sp['completed_qty'] = completed
                sp['remaining_qty'] = remaining
                if parent_static.get('time_study'):
                    # 1-operator hours for this split's quantity (see read_pdos).
                    sp['finished_hrs'] = round(s['total_qty'] / parent_static['time_study'], 2)
                    sp['std_hours']    = sp['finished_hrs']
                merged.append(sp)

        leftover = global_remaining - active_remaining_sum
        if parent_id in by_id:
            leftover_station, leftover_bin = find_pdo_current_location(parent_id, assignments)
            if leftover_bin == 'inprogress':
                leftover = max(0, leftover - get_wip_deduction(parent_id, leftover_station))
            if leftover <= 0:
                merged.remove(by_id[parent_id])
            else:
                by_id[parent_id]['remaining_qty'] = leftover
                by_id[parent_id]['completed_qty'] = parent_qty - leftover
                by_id[parent_id]['is_split_parent'] = True

    return merged

_PRIORITY_RANK = {'ph': 0, 'pm': 1, 'pl': 2}

def priority_bucket(p):
    """Mirrors the frontend's pc() classification exactly — needed as an
    explicit sort key because the day-threshold and urgency-threshold checks
    are independent ORs. Sorting by raw urgency alone doesn't guarantee
    High/Mid/Low grouping: a High card qualifying via days_left<=1 can still
    have a HIGHER urgency number than a Mid card qualifying via days_left<=6,
    since std_hours pulls urgency down independently of which day-bucket
    a card falls into.
    A GL-set manual override (ph/pm/pl) wins over the computed bucket.
    HIGH/red is reserved for orders due within 1 day (days_left<=1); the urgency
    score no longer promotes to High (it still surfaces early orders as Mid)."""
    ov = p.get('priority_override')
    if ov in _PRIORITY_RANK:
        return _PRIORITY_RANK[ov]
    if p['days_left'] <= 1:
        return 0  # High
    if p['days_left'] <= 6 or p['urgency'] < 35:
        return 1  # Mid
    return 2  # Low

def rebuild_final_pdo_cache():
    """Merge raw Excel PDOs with live SQL FG completion, filter out fully-completed
    orders, and sort by urgency. This is the function that produces what the
    board actually serves — called after either an Excel refresh OR an FG poll.
    Also folds in end-of-shift WIP (Nhaplecuoingay_All) for whatever's
    currently in-progress on a station — see get_wip_deduction()."""
    global _pdo_cache
    assignments = load_assignments()
    merged = []
    for p in _raw_pdo_cache:
        completed = _fg_completed_cache.get(p['id'], 0)
        completed = max(0, min(completed, p['qty']))
        # Partial-box WIP is real completed work on the PO, so count it wherever
        # it was last logged — not only while the order sits on its original
        # station. (Splits keep their own per-split, per-station WIP accounting
        # in inject_splits_into_cache; this parent-level add is overwritten there
        # for any order that has active splits.)
        completed = min(p['qty'], completed + get_wip_deduction_latest(p['id']))
        remaining = p['qty'] - completed
        if remaining <= 0:
            continue  # fully completed — drop from active list
        p2 = dict(p)
        p2['completed_qty'] = completed
        p2['remaining_qty'] = remaining
        p2['is_split_parent'] = False
        p2['is_split_child']  = False
        p2['split_letter']    = None
        p2['split_station']   = None
        p2['parent_id']       = None
        merged.append(p2)
    merged = inject_splits_into_cache(merged, assignments)
    # Attach any GL manual priority override (keyed by the PDO id, or the parent
    # id for split children) so it drives both the tag shown and the sort below.
    overrides = load_json(f'priority_{AREA}.json', {})
    if overrides:
        for x in merged:
            ov = overrides.get(x['id']) or overrides.get(x.get('parent_id') or '')
            x['priority_override'] = ov if ov in _PRIORITY_RANK else None
    else:
        for x in merged:
            x['priority_override'] = None
    # Received orders first (grouped), then priority bucket (High/Mid/Low) as
    # a HARD grouping, urgency only breaks ties WITHIN a bucket. See
    # priority_bucket() docstring for why urgency alone can't do this.
    merged.sort(key=lambda x: (0 if x['received'] > 0 else 1, priority_bucket(x), x['urgency']))
    _pdo_cache = merged

def _find_attendance_sheet_for_date(wb, target_date):
    """Auto-detect which monthly sheet contains target_date. Ported from server.py."""
    target_ym = (target_date.year, target_date.month)
    def extract_ym(sheet_name):
        m = re.search(r'T(\d{1,2})-(\d{4})', sheet_name)
        return (int(m.group(2)), int(m.group(1))) if m else None
    name_matches = [s for s in wb.sheetnames if extract_ym(s) == target_ym]
    for sname in reversed(name_matches):
        ws = wb[sname]
        rows = list(ws.iter_rows(max_row=2, values_only=True))
        if len(rows) > 1 and any(isinstance(v, datetime) and v.date() == target_date for v in rows[1]):
            return sname, ws
    for sname in reversed(wb.sheetnames):
        ws = wb[sname]
        rows = list(ws.iter_rows(max_row=2, values_only=True))
        if len(rows) > 1 and any(isinstance(v, datetime) and v.date() == target_date for v in rows[1]):
            return sname, ws
    return None, None

# Nominal full-day hours as literally recorded in the attendance sheet's OT
# column on a weekend (see _worker_window() below for why weekends are
# special) — this is the round number the sheet uses (8), not the app's
# precise shift-length decimals (8.167 etc.).
WEEKEND_BASE_SHIFT_HOURS = 8.0

def _worker_window(shift_name, ot_clock):
    """Given a worker's shift and today's OT hours, return (start, end) as
    decimal hours defining their actual working window today.

    Weekday: ot_clock is pure OT on top of an already-worked normal shift —
    buckets into 0 / ~2h / ~3h tiers matching the shift capacity summary's
    standard levels, same as always.

    Weekend: there is no separate "normal shift" baseline — every hour
    worked on a Sat/Sun gets logged in the same OT column, so the raw value
    means something different there:
      - >= 8h  → a full normal-length shift was worked; anything beyond 8h
                 is real OT, bucketed with the same tiers as a weekday
      - < 8h   → a partial day — the value IS the hours actually worked,
                 starting at the shift's normal start time and ending early
                 (e.g. a bare '5' means 5 of the normal ~8h, not 5h of OT)
    The attendance status column (ĐL/ĐC/HC) is read identically on weekends
    as weekdays — only the OT-column interpretation differs here."""
    s = SHIFTS.get(shift_name, SHIFTS['Ca 1'])
    if date.today().weekday() >= 5:  # Mon=0 ... Sat=5, Sun=6
        if ot_clock < WEEKEND_BASE_SHIFT_HOURS:
            return s['start'], s['start'] + ot_clock
        ot_clock = ot_clock - WEEKEND_BASE_SHIFT_HOURS  # excess beyond a full day = real OT
    if ot_clock >= 2.5:
        return s.get('ot3_start', s['start']), s.get('ot3_end', s['end'])
    elif ot_clock >= 1:
        return s.get('ot2_start', s['start']), s.get('ot2_end', s['end'])
    else:
        return s['start'], s['end']

# ── Break schedule & net-productive-hours model ────────────────────────────────
# From the official "Bảng tổng hợp thời gian làm việc theo ca" (shift capacity
# summary). Per (shift, OT tier): the clock start/end and every scheduled break
# inside it. Net productive hours = (shift_length − breaks − 5S) × (1 − PFD),
# where 5S is a fixed 5 min at the very start + 5 min at the very end of the
# shift (planned downtime) and PFD is a 6% personal-fatigue-delay allowance on
# the gross working time. For a FULL shift this reproduces the published
# figures (Ca1/Ca2: 6.73 / 8.47 / 9.08h; HC: 6.90 / 8.47h). It's used to value
# PARTIAL windows (half days / part-day lends) by netting only the breaks that
# fall inside the actual window worked. Full-shift, un-noted workers keep the
# exact published lookup (CAPACITY_HOURS_BY_TIER) instead.
def _hm(h, m):
    return h + m / 60.0

FIVE_S_MIN = 5.0            # minutes at each end of the shift (planned 5S downtime)
PFD_RATE   = 0.06          # 6% personal-fatigue-delay on gross working time

# (shift, tier) → (shift_start, shift_end, [ (break_start, break_end), ... ])
SHIFT_SCHEDULE = {
    ('Ca 1', 0): (_hm(6,0),  _hm(14,10), [(_hm(8,0),_hm(8,10)), (_hm(10,30),_hm(11,0)), (_hm(12,10),_hm(12,20))]),
    ('Ca 1', 2): (_hm(6,0),  _hm(16,10), [(_hm(8,0),_hm(8,10)), (_hm(10,30),_hm(11,0)), (_hm(12,10),_hm(12,20)), (_hm(14,10),_hm(14,20))]),
    ('Ca 1', 3): (_hm(6,0),  _hm(17,10), [(_hm(8,0),_hm(8,10)), (_hm(10,30),_hm(11,10)), (_hm(14,10),_hm(14,40))]),
    ('Ca 2', 0): (_hm(14,10),_hm(22,20), [(_hm(16,10),_hm(16,20)), (_hm(19,0),_hm(19,30)), (_hm(20,30),_hm(20,40))]),
    ('Ca 2', 2): (_hm(12,10),_hm(22,20), [(_hm(14,10),_hm(14,20)), (_hm(16,10),_hm(16,20)), (_hm(19,0),_hm(19,30)), (_hm(20,30),_hm(20,40))]),
    ('Ca 2', 3): (_hm(11,10),_hm(22,20), [(_hm(14,10),_hm(14,40)), (_hm(17,10),_hm(17,20)), (_hm(19,0),_hm(19,30)), (_hm(20,30),_hm(20,40))]),
    ('HC',   0): (_hm(8,0),  _hm(16,30), [(_hm(10,0),_hm(10,10)), (_hm(12,0),_hm(12,40)), (_hm(15,0),_hm(15,10))]),
    ('HC',   2): (_hm(8,0),  _hm(18,30), [(_hm(10,0),_hm(10,10)), (_hm(11,0),_hm(11,10)), (_hm(12,0),_hm(12,40)), (_hm(15,0),_hm(15,10)), (_hm(16,30),_hm(16,40))]),
    # HC has no 3h-OT row in the summary; _shift_downtime() falls back to HC/OT2.
}

def _shift_downtime(shift, tier):
    """(shift_start, shift_end, downtime_intervals) for a (shift, tier), where
    downtime = the scheduled breaks plus the two 5S blocks at each shift end."""
    sched = SHIFT_SCHEDULE.get((shift, tier)) or SHIFT_SCHEDULE.get((shift, 2)) or SHIFT_SCHEDULE.get((shift, 0)) or SHIFT_SCHEDULE[('Ca 1', 0)]
    s_start, s_end, breaks = sched
    five = FIVE_S_MIN / 60.0
    downtime = list(breaks) + [(s_start, s_start + five), (s_end - five, s_end)]
    return s_start, s_end, downtime

def net_hours_for_windows(shift, tier, windows):
    """Net productive hours across one or more actual working windows (decimal
    hours), netting out every break / 5S block that overlaps each window and
    applying the 6% PFD to the remaining gross time. Windows are clamped to the
    shift's own clock bounds."""
    s_start, s_end, downtime = _shift_downtime(shift, tier)
    total_net_min = 0.0
    for (ws, we) in windows:
        ws = max(ws, s_start); we = min(we, s_end)
        if we <= ws:
            continue
        clock_min = (we - ws) * 60.0
        dt_min = 0.0
        for (bs, be) in downtime:
            ov = min(we, be) - max(ws, bs)
            if ov > 0:
                dt_min += ov * 60.0
        gross = max(0.0, clock_min - dt_min)
        total_net_min += gross * (1 - PFD_RATE)
    return round(total_net_min / 60.0, 2)

# Time-window note parser. A red-corner Excel comment on a worker's status cell
# can spell out the actual hours worked and where, e.g.:
#   "6:00-11:00"                     → half day, in THIS area
#   "RP 6:00-11:00, PK 11:00-14:10"  → in RP until 11:00, then lent to PK
# For the app instance of a given AREA we keep only the segments tagged with
# that area (or untagged = this area); segments tagged for another area mean the
# worker is elsewhere then. Times accept H:MM / HH:MM / H.MM / HHhMM and a
# plain '-', en-dash or em-dash separator.
_AREA_CODES = ('PK', 'AD', 'RP')
_NOTE_SEG_RE = re.compile(
    r'([A-Za-z]{2,})?\s*(\d{1,2})\s*[:h.]\s*(\d{2})\s*[-–—]\s*(\d{1,2})\s*[:h.]\s*(\d{2})',
    re.IGNORECASE)

def parse_area_note(text, area):
    """Parse an attendance time-window note for THIS area.
    Returns {'parsed': False} if no time range is found, else
    {'parsed': True, 'in_area': [(s,e)...], 'other': [(code,s,e)...]}."""
    if not text:
        return {'parsed': False}
    in_area, other = [], []
    for m in _NOTE_SEG_RE.finditer(str(text)):
        code = (m.group(1) or '').upper()
        s = int(m.group(2)) + int(m.group(3)) / 60.0
        e = int(m.group(4)) + int(m.group(5)) / 60.0
        if e <= s or s < 0 or e > 24:
            continue
        if code in _AREA_CODES:
            if code == area:
                in_area.append((s, e))
            else:
                other.append((code, s, e))
        else:
            in_area.append((s, e))   # untagged (or a non-area word) → this area
    if not in_area and not other:
        return {'parsed': False}
    return {'parsed': True, 'in_area': in_area, 'other': other}

def _cell_note_text(cell):
    """The red-corner comment text on a cell, or '' — tolerant of None cell /
    no comment. openpyxl only populates .comment when NOT in read-only mode."""
    try:
        c = getattr(cell, 'comment', None)
        return (c.text or '').strip() if c is not None else ''
    except Exception:
        return ''

def _shift_from_text(text):
    """Leading shift code found in any whitespace/comma/semicolon token of text
    → 'Ca 1' (ĐL) / 'Ca 2' (ĐC) / 'HC', else None. Works for both the status
    cell and a note like 'ĐL 6:00-11:00'."""
    if not text:
        return None
    for tok in re.split(r'[\s,;]+', str(text).upper()):
        if tok.startswith('ĐL'):
            return 'Ca 1'
        if tok.startswith('ĐC'):
            return 'Ca 2'
        if tok.startswith('HC'):
            return 'HC'
    return None

def _find_date_col(rows, target_date):
    """Index of the header column (row 2) holding target_date, or None."""
    if len(rows) < 2:
        return None
    for idx, cell in enumerate(rows[1]):
        v = cell.value
        if isinstance(v, datetime) and v.date() == target_date:
            return idx
    return None

def _workers_from_rows(rows, date_col, area, include_lent_away=False):
    """Parse the per-worker attendance for one date column out of already-loaded
    sheet rows (list of Cell tuples). Factored out of read_attendance_shift_info
    so it can be reused for a single day OR every day in a range (capacity chart)
    without re-reading the workbook. See read_attendance_shift_info for the
    row-layout and lend/borrow/note rules."""
    def cval(cells, idx):
        return cells[idx].value if idx < len(cells) else None

    data_rows = rows[2:]
    out = []
    i = 0
    while i + 2 < len(data_rows):
        main_cells, row2_cells, row3_cells = data_rows[i], data_rows[i+1], data_rows[i+2]
        c0 = cval(main_cells, 0)
        if not isinstance(c0, (int, float)) or c0 != int(c0):
            i += 1
            continue
        msnv = str(cval(main_cells, 2)).strip() if cval(main_cells, 2) else ""
        if msnv in INDIRECT_LABOR_IDS:
            i += 3
            continue

        # Row layout per worker:
        #   ORIGIN worker   → [status, OT,     Lend]   (row2 A="OT",    row3 A="Lend")
        #   BORROWED worker → [status, Borrow, OT]     (row2 A="Borrow", row3 A="OT")
        # So the OT figure is always in whichever row is the OT row, and the
        # Lend row only exists for origin workers.
        row2_label = str(cval(row2_cells, 0)).strip().lower() if cval(row2_cells, 0) else ""
        row3_label = str(cval(row3_cells, 0)).strip().lower() if cval(row3_cells, 0) else ""
        is_borrow = (row2_label == "borrow")
        ot_cells = row3_cells if is_borrow else row2_cells

        status = cval(main_cells, date_col)
        status_str = str(status).strip() if status is not None else ""

        # Time-window note (red-corner comment): check the status cell first, then
        # the OT / lend row's date cell as a fallback for wherever it was typed.
        note_text = (_cell_note_text(main_cells[date_col] if date_col < len(main_cells) else None)
                     or _cell_note_text(row2_cells[date_col] if date_col < len(row2_cells) else None)
                     or _cell_note_text(row3_cells[date_col] if date_col < len(row3_cells) else None))
        note = parse_area_note(note_text, area)
        note_has_in_area = bool(note.get('parsed') and note.get('in_area'))

        # Shift comes from the LEADING code in the status cell (ĐL→Ca1, ĐC→Ca2,
        # HC). Absent/half-day-off codes (NP 1/2, NĐB, ...) yield no code.
        shift = _shift_from_text(status_str)
        infer = False
        if not shift:
            # Off/absent code. Normally skipped — BUT a half-day worker may be
            # marked with an off-code (e.g. 'NP 1/2') yet still work part of the
            # day, recorded via a time-window note. If that note has an in-area
            # window, honour it: keep the worker, taking the shift from a code in
            # the note (e.g. 'ĐL 6:00-11:00'). With no code in the note, the shift
            # is resolved AFTER the loop from the work time (see the post-pass) —
            # the placeholder here is replaced then.
            if not note_has_in_area:
                i += 3
                continue  # genuinely absent today — not available
            shift = _shift_from_text(note_text)
            if not shift:
                infer = True
                shift = 'Ca 1'   # placeholder — real shift decided in the post-pass

        # LEND (origin workers only): the numeric "Lend" row. A FULL-shift lend
        # (|v| ≥ ~8h) means the worker is in ANOTHER area all day; a partial lend
        # keeps them here. A time-window note (below) supersedes this when present.
        lend_hours = 0.0
        if not is_borrow and row3_label == "lend":
            lv = cval(row3_cells, date_col)
            if isinstance(lv, (int, float)):
                lend_hours = float(lv)
        lent_away = (not is_borrow) and abs(lend_hours) >= WEEKEND_BASE_SHIFT_HOURS

        ot_clock = 0
        ot_val = cval(ot_cells, date_col)
        if isinstance(ot_val, (int, float)) and ot_val > 0:
            ot_clock = float(ot_val)
        tier = _ot_tier(ot_clock)

        note_kind = ''
        windows = []
        if note.get('parsed'):
            if note['in_area']:
                windows   = note['in_area']
                w_start   = min(s for s, e in windows)
                w_end     = max(e for s, e in windows)
                cap_hours = net_hours_for_windows(shift, tier, windows)
                note_kind = 'partial'
            else:
                # Note present but every segment is in ANOTHER area → away today.
                w_start, w_end = _worker_window(shift, ot_clock)
                cap_hours = 0.0
                note_kind = 'away'
                lent_away = True
        else:
            w_start, w_end = _worker_window(shift, ot_clock)
            cap_hours = operator_capacity_hours(shift, ot_clock)

        if lent_away and note_kind == '':
            cap_hours = 0.0   # numeric full-shift lend, no overriding note

        if lent_away and not include_lent_away:
            i += 3
            continue  # not available in this area's pool/capacity today

        entry = {
            'id':         msnv,
            'name':       str(cval(main_cells, 1) or ''),
            'shift':      shift,
            'ot_clock':   ot_clock,
            'is_borrow':  is_borrow,      # borrowed IN from another area
            'lend_hours': lend_hours,     # negative = lent out (partial if |v| < full shift)
            'lent_away':  lent_away,      # full shift elsewhere → excluded from pool/capacity
            'start':      w_start,        # effective window start (note-derived if noted)
            'end':        w_end,          # effective window end
            'cap_hours':  round(cap_hours, 2),   # productive hours (net-of-breaks if noted)
            'windows':    windows,        # in-area note segments (for the gantt), [] otherwise
            'note':       note_text,      # raw note text for display
            'note_kind':  note_kind,      # '' normal · 'partial' in-area window · 'away' elsewhere
        }
        if infer:
            # Shift not stated (off-code cell + note with no shift code) — carry
            # what's needed to decide it in the post-pass below.
            entry['_infer'] = True
            entry['_infer_tier'] = tier
        out.append(entry)
        i += 3

    # Resolve inferred shifts. Match the actual work window to the FIRST (Ca 1)
    # or SECOND (Ca 2) shift by its start time; only fall back to HC when the day
    # is predominantly a general-shift day (≥80% of the coded operators are HC).
    coded = [w for w in out if not w.get('_infer')]
    hc_n = sum(1 for w in coded if w['shift'] == 'HC')
    is_hc_day = bool(coded) and (hc_n / len(coded)) >= 0.8
    for w in out:
        if w.pop('_infer', False):
            tier = w.pop('_infer_tier', 0)
            # Use _hm(14,10) — computed the SAME way the note time is (h + m/60)
            # — as the Ca1/Ca2 boundary. CA1_END is the rounded literal 14.167,
            # so a window starting exactly at 14:10 (parsed 14.16667) compared
            # against it read as "before Ca1 end" and mis-tagged the worker Ca 1
            # (blue). A 14:10 start is the Ca 2 boundary → afternoon shift.
            w['shift'] = 'HC' if is_hc_day else ('Ca 1' if w['start'] < _hm(14, 10) else 'Ca 2')
            if w['windows']:
                w['cap_hours'] = round(net_hours_for_windows(w['shift'], tier, w['windows']), 2)

    return out

def read_attendance_shift_info(area, target_date, include_lent_away=False):
    """Lightweight attendance parse — returns per-worker shift + OT bucket,
    plus lend/borrow state (is_borrow, lend_hours, lent_away).
    Row format ported from server.py's parse_capacity_file: each worker has
    3 rows (status / OT / lend-borrow), keyed by MSNV in column index 2.

    A FULL-shift lend means the worker is physically in another area all day.
    By default such workers are dropped (not in this area's pool/capacity); pass
    include_lent_away=True to keep them in the list flagged lent_away=True so the
    attendance report can show them (with a note) even though they stay
    unassignable. Borrowed-in workers are always included, flagged is_borrow."""
    # Track whether TODAY's read succeeded so prune_expired_workers can refuse to
    # wipe operators on a failed load. Pessimistic: assume failure until we reach a
    # clean parse below. (Only today's read gates the prune; other dates don't.)
    global _last_attendance_ok
    _track_today = (target_date == date.today())
    if _track_today:
        _last_attendance_ok = False

    filepath = CAPACITY_FILES.get(area)
    if not filepath:
        return []
    path = filepath if os.path.exists(filepath) else DEV_ATTENDANCE
    if not os.path.exists(path):
        return []

    tmp_path = None
    try:
        # Full (non-read-only) load so red-corner cell COMMENTS are available —
        # openpyxl only populates cell.comment outside read-only mode. Heavier
        # than the streaming read, but the file is ~1MB and this runs on the
        # 5-minute worker-refresh throttle. Still read from a temp copy and never
        # saved back, so there is no risk to the source file.
        wb, tmp_path = _safe_load_workbook(path, read_only=False, data_only=True)
        sheet_name, ws = _find_attendance_sheet_for_date(wb, target_date)
        if ws is None:
            wb.close()
            return []
        rows = list(ws.iter_rows())   # Cell objects — need both .value and .comment
        wb.close()
    except Exception as e:
        print(f'[Attendance error] {e}')
        return []
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try: os.remove(tmp_path)
            except: pass

    date_col = _find_date_col(rows, target_date)
    if date_col is None:
        return []
    result = _workers_from_rows(rows, date_col, area, include_lent_away)
    if _track_today:
        _last_attendance_ok = True   # file + today's sheet/column all read cleanly
    return result

def compute_capacity_by_day(area, from_date, to_date):
    """Available productive labour-hours for EACH day in [from_date, to_date],
    read from that day's own attendance column. Returns {YYYY-MM-DD: hours}, and
    ONLY includes days that actually have attendance filled (capacity > 0) — a
    day with no data is omitted so the chart can leave a gap there rather than
    plot a misleading zero. The workbook is loaded once; rows are cached per
    monthly sheet so a multi-week range is a single file read."""
    filepath = CAPACITY_FILES.get(area)
    if not filepath:
        return {}
    path = filepath if os.path.exists(filepath) else DEV_ATTENDANCE
    if not os.path.exists(path):
        return {}
    result = {}
    tmp_path = None
    try:
        wb, tmp_path = _safe_load_workbook(path, read_only=False, data_only=True)
        rows_cache = {}   # sheet_name → loaded rows (avoid re-reading a month per day)
        d = from_date
        while d <= to_date:
            sheet_name, ws = _find_attendance_sheet_for_date(wb, d)
            if ws is not None:
                if sheet_name not in rows_cache:
                    rows_cache[sheet_name] = list(ws.iter_rows())
                rows = rows_cache[sheet_name]
                date_col = _find_date_col(rows, d)
                if date_col is not None:
                    workers = _workers_from_rows(rows, date_col, area, include_lent_away=False)
                    total = round(sum(w['cap_hours'] for w in workers), 1)
                    if total > 0:
                        result[d.strftime('%Y-%m-%d')] = total
            d += timedelta(days=1)
        wb.close()
    except Exception as e:
        print(f'[capacity_by_day error] {e}')
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try: os.remove(tmp_path)
            except: pass
    return result

def get_available_workers(area='PK'):
    """TODAY's full on-shift roster (everyone present per attendance, each with
    their working window). Membership by current time — the 15-min pre-shift
    pre-open, the 'active' flag, and shift-end removal — is applied LIVE in
    build_board_snapshot() / prune_expired_workers() via _worker_available_now(),
    NOT here. That keeps those transitions accurate to the minute even though the
    attendance file itself is only re-read on the ~5-min throttle (previously the
    pool membership was frozen between reads, so e.g. a worker didn't appear at
    05:50 for a 06:00 shift until the next file read)."""
    now = datetime.now()
    now_decimal = now.hour + now.minute / 60
    info = read_attendance_shift_info(area, date.today())
    roster = []
    for w in info:
        start, end = w.get('start'), w.get('end')
        if start is None or end is None:
            start, end = _worker_window(w['shift'], w['ot_clock'])
        roster.append({
            'id': w['id'], 'name': w['name'], 'shift': w['shift'],
            'start': start, 'end': end,
            'active': start <= now_decimal <= end,
            'note': w.get('note', ''), 'note_kind': w.get('note_kind', ''),
        })
    return roster

def _worker_available_now(w, now_decimal):
    """Whether a roster worker is currently selectable: from 15 min before their
    shift starts (pre-open) through the end of their window."""
    start, end = w.get('start'), w.get('end')
    if start is None or end is None:
        return bool(w.get('active', True))
    return (start - NEXT_SHIFT_OPEN_EARLY_HRS) <= now_decimal <= end

def _active_worker_id_set():
    """IDs of workers whose shift is actually underway RIGHT NOW. Computed live
    from each worker's stored window so the active/upcoming transition (which
    gates ETC) is accurate to the minute, not stale by up to the 5-min
    attendance-refresh interval."""
    now_decimal = datetime.now().hour + datetime.now().minute / 60
    ids = set()
    for w in _worker_cache:
        st, en = w.get('start'), w.get('end')
        if st is None or en is None:
            if w.get('active', True):
                ids.add(str(w['id']))
        elif st <= now_decimal <= en:
            ids.add(str(w['id']))
    return ids

def _count_active_workers(station, workers_asgn):
    """How many of the workers assigned to `station` are currently active — the
    count that should drive ETC and the paused state (a pre-assigned upcoming
    worker shouldn't make the clock tick until their shift starts)."""
    active = _active_worker_id_set()
    return sum(1 for wid in workers_asgn.get(station, []) if str(wid) in active)

# Productive labour-hours a single present operator contributes, by shift
# category and OT tier — net of breaks, per production-planning's figures
# (NOT the raw clock window length, which includes unpaid breaks). Ca1 and Ca2
# share one column; the general (HC) shift differs at the no-OT tier.
#   tier 0 = no OT · tier 2 = ~2h OT · tier 3 = ~3h OT
# HC + 3h OT was not separately specified, so it mirrors the Ca 3h-OT value
# (HC and Ca already coincide at the 2h-OT tier); adjust HC[3] if that differs.
CAPACITY_HOURS_BY_TIER = {
    'CA': {0: 6.73, 2: 8.47, 3: 9.08},
    'HC': {0: 6.90, 2: 8.47, 3: 9.08},
}

def _ot_tier(ot_clock):
    """Bucket the attendance OT-column value into a productive-hours tier.
    Same thresholds as _worker_window so the gantt window and the capacity
    figure always agree on whether a worker is on 2h vs 3h OT."""
    if ot_clock >= 2.5:
        return 3
    if ot_clock >= 1:
        return 2
    return 0

def operator_capacity_hours(shift_name, ot_clock):
    """Productive hours one present operator contributes today, from their
    shift + OT tier."""
    cat = 'HC' if shift_name == 'HC' else 'CA'
    return CAPACITY_HOURS_BY_TIER[cat][_ot_tier(ot_clock)]

def compute_day_capacity_hours(area='PK'):
    """Total available productive labour-hours for today = sum of each present
    operator's capacity hours. Drives the capacity reference line on the report's
    hours chart. Fully-lent-away operators are already excluded by
    read_attendance_shift_info; borrowed-in operators are included. Each worker's
    cap_hours is the full-shift tier figure, OR — when a time-window note is
    present (half day / part-day lend) — the net-of-breaks hours for the actual
    window(s) worked in this area."""
    info = read_attendance_shift_info(area, date.today())
    total = 0.0
    for w in info:
        total += w.get('cap_hours', operator_capacity_hours(w['shift'], w['ot_clock']))
    return round(total, 1)

def compute_shift_windows(area='PK'):
    """Aggregate TODAY's actual attendance into effective start/end times per
    shift, so the day-arc widget shows the real overlap window (driven by who's
    actually on OT today) rather than just the theoretical 10-minute handoff.
    Falls back to the static SHIFTS config if attendance data isn't reachable."""
    info = read_attendance_shift_info(area, date.today())
    fallback = {
        'mode': 'split',
        'ca1_start': SHIFTS['Ca 1']['start'], 'ca1_end': SHIFTS['Ca 1']['end'],
        'ca2_start': SHIFTS['Ca 2']['start'], 'ca2_end': SHIFTS['Ca 2']['end'],
        'source': 'fallback_static',
    }
    if not info:
        return fallback

    ca1 = [w for w in info if w['shift'] == 'Ca 1']
    ca2 = [w for w in info if w['shift'] == 'Ca 2']
    hc  = [w for w in info if w['shift'] == 'HC']

    def agg_window(workers, shift_name):
        if not workers:
            return None
        # Use each worker's effective window (note-derived if any), falling back
        # to the shift window.
        windows = [(w.get('start'), w.get('end')) for w in workers]
        windows = [wn if wn[0] is not None else _worker_window(shift_name, w['ot_clock'])
                   for wn, w in zip(windows, workers)]
        return min(s for s, e in windows), max(e for s, e in windows)

    if not ca1 and not ca2 and hc:
        # General-shift-only day — single HC zone, no Ca1/Ca2 split
        win = agg_window(hc, 'HC')
        return {'mode': 'hc', 'hc_start': win[0], 'hc_end': win[1], 'source': 'attendance'}

    ca1_win = agg_window(ca1, 'Ca 1')
    ca2_win = agg_window(ca2, 'Ca 2')
    if not ca1_win and not ca2_win:
        return fallback  # no one on either shift today (odd, but don't show a broken widget)

    ca1_start, ca1_end = ca1_win or (SHIFTS['Ca 1']['start'], SHIFTS['Ca 1']['start'])
    ca2_start, ca2_end = ca2_win or (SHIFTS['Ca 2']['end'],   SHIFTS['Ca 2']['end'])
    return {
        'mode': 'split',
        'ca1_start': ca1_start, 'ca1_end': ca1_end,
        'ca2_start': ca2_start, 'ca2_end': ca2_end,
        'source': 'attendance',
    }

def stations_path():
    return _path(f'stations_{AREA}.json')

def load_stations():
    """Stations are now managed entirely in the web UI (add/remove), persisted
    as JSON — database.xlsx's Station sheet is retired, it had no other purpose."""
    p = stations_path()
    if not os.path.exists(p):
        fallback = get_fallback_stations()
        save_stations(fallback)  # seed on first run
        return list(fallback)
    try:
        with open(p, encoding='utf-8') as f:
            stations = json.load(f)
        return stations if stations else get_fallback_stations()
    except Exception as e:
        print(f'[stations load error] {e}')
        return get_fallback_stations()

def save_stations(stations):
    with open(stations_path(), 'w', encoding='utf-8') as f:
        json.dump(stations, f, indent=2)

def refresh_cache(force=False):
    global _raw_pdo_cache, _pdo_cache, _worker_cache, _station_cache, _last_excel_mtime, _last_pdo_status
    global _shift_windows_cache, _last_worker_refresh_ts, _day_capacity_hours
    path = _resolve_path(EXCEL_PLAN, DEV_PLAN)
    try:
        mtime = os.path.getmtime(path)
    except:
        mtime = 0
    with _cache_lock:
        try:
            if force or mtime != _last_excel_mtime or not _raw_pdo_cache:
                raw, status = read_pdos()
                _last_pdo_status = status
                if status == 'ok':
                    _raw_pdo_cache = raw
                    _last_excel_mtime = mtime
                elif status == 'invalid_filename':
                    # Deliberate hard block, not a transient error — clear the
                    # cache rather than keep serving stale/wrong-area data.
                    _raw_pdo_cache = []
                rebuild_final_pdo_cache()  # re-merge with whatever FG data we have, even on error (keeps old data visible)
        except Exception as e:
            print(f'[refresh_cache PDO error] {e}')
            _last_pdo_status = 'error'
        try:
            # Worker pool + shift windows both read the attendance file (relatively
            # heavy — file copy + openpyxl parse). Throttled to avoid re-reading it
            # on every ~10s board poll; still refreshes on force (e.g. startup).
            now_ts = time.time()
            if force or (now_ts - _last_worker_refresh_ts) > WORKER_REFRESH_THROTTLE_SEC:
                _worker_cache = get_available_workers(AREA)
                _shift_windows_cache = compute_shift_windows(AREA)
                _day_capacity_hours = compute_day_capacity_hours(AREA)
                _last_worker_refresh_ts = now_ts
        except Exception as e:
            print(f'[refresh_cache worker error] {e}')
        try:
            if not _station_cache:
                stations = load_stations()
                if stations:
                    _station_cache = stations
                else:
                    _station_cache = get_fallback_stations()
        except Exception as e:
            print(f'[refresh_cache station error] {e}')
            if not _station_cache:
                _station_cache = get_fallback_stations()

def refresh_fg_and_advance():
    """The 1-minute poll cycle: pull fresh FG completion from SQL, re-merge into
    the PDO cache, auto-advance any station whose in-progress order just finished,
    drop expired-shift workers, AND check the plan Excel file for changes (mtime)
    so it's picked up automatically without needing a manual 'Refresh orders' click."""
    global _fg_completed_cache, _wip_cache, _fg_latest_scan_cache, _license_plate_cache
    # Per-phase timing so the server console shows exactly where a slow refresh
    # goes (Excel read vs. each SQL query vs. splits) — invaluable on the plant
    # network where any one of these can dominate.
    _t = {}
    _mark0 = time.perf_counter()
    def _lap(name, since):
        _t[name] = time.perf_counter() - since
        return time.perf_counter()

    _c = _mark0
    try:
        refresh_cache()  # non-forced — only actually re-reads Excel if its mtime changed
    except Exception as e:
        print(f'[plan file follow error] {e}')
    _c = _lap('excel', _c)

    try:
        # Restrict the FG SUM to just the POs on the board (huge win vs. an
        # all-time full-table GROUP BY on a large FG log).
        po_ids = [p['id'] for p in _raw_pdo_cache]
        data = get_fg_completed(po_ids)
        _c = _lap('fg', _c)
        wip_data = get_wip_entries()
        _c = _lap('wip', _c)
        scan_data = get_fg_latest_scan_times()
        _c = _lap('scan', _c)
        lp_data = get_license_plate_prints()  # cached so the report can show LP tags too
        with _cache_lock:
            if data or _last_fg_status == 'ok':
                _fg_completed_cache = data
            # Only overwrite on a successful query (None = the SQL call failed);
            # a transient error must not zero out WIP/scan deductions and make
            # remaining_qty flicker upward for a cycle.
            if wip_data is not None:
                _wip_cache = wip_data
            if scan_data is not None:
                _fg_latest_scan_cache = scan_data
            if lp_data is not None:
                _license_plate_cache = lp_data
            else:
                # A None return means the License Plate SQL call FAILED (see the
                # [License Plate SQL error] line just above). We keep the last-good
                # cache so tags don't flicker — but that also means any order
                # printed AFTER the last good fetch is invisible to auto-queue and
                # will NOT auto-place until the query recovers. Make that loud:
                # a silently-stale LP cache looks exactly like "auto-queue is
                # broken" from the floor.
                print('[License Plate] query FAILED — auto-queue is running on STALE '
                      'print data; orders printed since the last good fetch will NOT '
                      'auto-place until this recovers.')
            refresh_split_completions()
            _c = _lap('splits', _c)
            rebuild_final_pdo_cache()
            _c = _lap('rebuild', _c)
    except Exception as e:
        print(f'[FG poll error] {e}')

    try:
        check_auto_advance_from_fg()
    except Exception as e:
        print(f'[auto-advance error] {e}')

    # Authoritative prune — now that the cache was just fully rebuilt from fresh
    # Excel+SQL, clear any assignment references to PDOs that are genuinely gone.
    # This is the ONLY place pruning happens now (build_board_snapshot no longer
    # does it on every GET), so it acts on a consistent, freshly-built view.
    try:
        with _cache_lock:
            _asg = load_assignments()
            _tmr = load_timers()
            if prune_stale_assignments(_asg, _tmr):
                save_assignments(_asg)
                save_timers(_tmr)
    except Exception as e:
        print(f'[prune error] {e}')

    try:
        auto_queue_from_license_plate()
    except Exception as e:
        print(f'[auto-queue error] {e}')
    _c = _lap('license_plate+advance', _c)

    try:
        reset_daily_worker_staffing()   # once/day: drop carried-over staffing from a missed shift-end
    except Exception as e:
        print(f'[daily reset error] {e}')

    try:
        prune_expired_workers()
    except Exception as e:
        print(f'[worker prune error] {e}')

    # Shift-boundary ETC reset: when we cross into a new shift phase (Ca1 end →
    # Ca2, or overnight → morning Ca1), reset every station's ETC so it's
    # recalculated from the current remaining qty. First poll just records the
    # phase (no spurious reset on startup).
    try:
        global _last_shift_phase
        phase = current_shift_phase()
        if _last_shift_phase is None:
            _last_shift_phase = phase
        elif phase != _last_shift_phase:
            prev = _last_shift_phase
            _last_shift_phase = phase
            print(f'[shift phase] {prev} → {phase}')
            if phase in ('ca1', 'ca2'):  # entering a production shift → fresh ETC
                reset_all_timers_for_new_shift()
    except Exception as e:
        print(f'[shift phase error] {e}')

    total = time.perf_counter() - _mark0
    if total >= 1.0:  # don't spam the console for fast no-op cycles
        breakdown = ' | '.join(f'{k} {v:.1f}s' for k, v in _t.items() if v >= 0.05)
        print(f'[refresh timing] total {total:.1f}s — {breakdown}')

def _fg_poll_loop():
    while True:
        time.sleep(FG_POLL_INTERVAL_SEC)
        refresh_fg_and_advance()

# ── ETC helpers ────────────────────────────────────────────────────────────────
def calc_etc(finished_hrs, remaining_qty, total_qty, worker_count, started_at_iso):
    """Return ISO timestamp of estimated completion for the CURRENT order only."""
    if not finished_hrs or not total_qty or not worker_count or not started_at_iso:
        return None
    hrs_per_unit = finished_hrs / total_qty
    remaining_hrs = hrs_per_unit * remaining_qty / max(worker_count, 1)
    started = datetime.fromisoformat(started_at_iso)
    etc = started.timestamp() + remaining_hrs * 3600
    return datetime.fromtimestamp(etc).isoformat()

def calc_queue_clear(ip_pdo, queue_pdos, worker_count, started_at_iso):
    """Return ISO timestamp when the ENTIRE station queue (current + all queued) clears.
    Returns None when unstaffed — an estimate assuming a phantom worker is misleading."""
    if not worker_count:
        return None
    total_hrs = 0.0
    # Remaining hours on current in-progress order (same math whether or not it
    # has a start timestamp — the timestamp only shifts the base time below).
    if ip_pdo:
        hrs_per_unit = (ip_pdo['finished_hrs'] / ip_pdo['qty']) if ip_pdo.get('qty') else 0
        total_hrs += hrs_per_unit * ip_pdo.get('remaining_qty', 0) / worker_count
    # Full hours for every queued order (not yet started, so full remaining_qty)
    for p in queue_pdos:
        if not p: continue
        hrs_per_unit = (p['finished_hrs'] / p['qty']) if p.get('qty') else 0
        total_hrs += hrs_per_unit * p.get('remaining_qty', 0) / worker_count
    if total_hrs <= 0:
        return None
    base_time = datetime.fromisoformat(started_at_iso) if started_at_iso else datetime.now()
    clear_ts = base_time.timestamp() + total_hrs * 3600
    return datetime.fromtimestamp(clear_ts).isoformat()

def pdo_by_id(pdo_id):
    # No lock here — callers that need thread safety hold the lock themselves
    return next((p for p in _pdo_cache if p['id'] == pdo_id), None)

def find_pdo_current_location(pdo_id, assignments):
    """Returns (station, bin_type) or (None, None) if not currently assigned anywhere."""
    for station, bins in assignments.items():
        if bins.get('inprogress') == pdo_id:
            return station, 'inprogress'
        if pdo_id in bins.get('queue', []):
            return station, 'queue'
    return None, None

def prune_stale_assignments(assignments, timers):
    """Clear any inprogress/queue references to PDOs that no longer exist in
    the live PDO cache (e.g. completed in Excel, or received reset to 0).
    Without this, a station can appear empty on the board while the server
    still thinks it's occupied — blocking all new assignments to it.

    GUARD: only prune when the plan read is trustworthy (_last_pdo_status == 'ok').
    On a failed/locked/unsynced read — especially a fresh startup where there's no
    last-good cache to fall back on — _pdo_cache is empty, so EVERY assigned PDO
    would look 'not in cache' and get cleared, wiping the whole board and saving the
    empty state. A transient read failure must never destroy operator-entered
    assignments; skip and try again next poll once the plan is readable again."""
    if _last_pdo_status != 'ok':
        print(f'[prune] skipped — plan status is {_last_pdo_status!r}, not \'ok\'; '
              f'refusing to prune assignments against an untrusted/empty cache')
        return False
    changed = False
    for station, bins in assignments.items():
        ip = bins.get('inprogress')
        if ip and pdo_by_id(ip) is None:
            # Logged so a removal is never silent — if orders ever leave a
            # station unexpectedly, the console shows which one and where.
            print(f'[prune] {station}: in-progress {ip} not in active cache — clearing '
                  f'(completed, or dropped by Excel filter / FG / WIP)')
            bins['inprogress'] = None
            timers.pop(station, None)
            changed = True
        q = bins.get('queue', [])
        cleaned_q = [pid for pid in q if pdo_by_id(pid) is not None]
        if len(cleaned_q) != len(q):
            dropped = [pid for pid in q if pid not in cleaned_q]
            print(f'[prune] {station}: removing queued {dropped} not in active cache')
            bins['queue'] = cleaned_q
            changed = True
    return changed

def check_auto_advance_from_fg():
    """Called after each FG poll. If a station's in-progress PDO has fully
    completed (per live SQL — meaning it dropped out of _pdo_cache), promote
    the next queued PDO into its place, same as the old manual FG-signal flow."""
    with _cache_lock:
        assignments = load_assignments()
        timers = load_timers()
        changed = False
        for station, bins in assignments.items():
            ip_id = bins.get('inprogress')
            if not ip_id:
                continue
            if pdo_by_id(ip_id) is None:  # completed — no longer in active cache
                q = bins.get('queue', [])
                if q:
                    new_ip = q.pop(0)
                    bins['inprogress'] = new_ip
                    timers[station] = {'started_at': datetime.now().isoformat(), 'fg_count': 0}
                    print(f'[auto-advance] {station}: {ip_id} completed → {new_ip} promoted')
                else:
                    bins['inprogress'] = None
                    timers.pop(station, None)
                    print(f'[auto-advance] {station}: {ip_id} completed → station now idle')
                changed = True
        if changed:
            save_assignments(assignments)
            save_timers(timers)

def auto_queue_from_license_plate():
    """A label-printed (License Plate) PDO that is active with work still
    remaining, and isn't currently on the board anywhere, gets auto-placed on
    its printed station. This is SELF-HEALING: it runs every 60s poll, so even
    if the board missed the exact print moment (app not polling, PC hiccup,
    OneDrive not synced, etc.) the order is recovered on a later poll.

    Placement is gated by four guards, and ONLY these four:
      1. active/unfinished — pdo_by_id(po) is not None (remaining_qty > 0;
         finished orders are already dropped from _pdo_cache, so they never
         come back).
      2. not already on the board — find_pdo_current_location finds nothing.
      3. not manually removed — suppression (lp_suppress_{AREA}.json): a GL
         removal blocks re-add until a NEWER LP print arrives.
      4. printed station is a real configured station on this board.

    It deliberately does NOT skip on existing FG/WIP activity. WIP
    (Nhaplecuoingay_All) and FG scans are logged by operators in SQL
    INDEPENDENTLY of the board's assignment state — an order can accumulate
    real progress while never having been placed on the board. Skipping those
    (the old behavior) permanently stranded any printed order that picked up
    WIP/FG before it was ever placed: the guard saw the activity and refused to
    auto-add it, so it could never appear. remaining_qty (which already nets
    out that WIP/FG) is the correct signal for "is there still work to do", and
    guard #1 covers it. Once a PDO lands in assignments, guard #2 stops it
    being re-added — no separate 'already processed' bookkeeping needed."""
    prints = _license_plate_cache  # populated by the poll (see refresh_fg_and_advance)
    if not prints:
        return
    with _cache_lock:
        assignments = load_assignments()
        station_ids = set(_station_cache) if _station_cache else set(get_fallback_stations())
        # PDOs the GL manually removed are suppressed until a NEWER LP print
        # arrives (compared by timestamp) — so a removed order doesn't keep
        # auto-jumping back into the queue on every poll.
        suppress = load_json(f'lp_suppress_{AREA}.json', {})
        # Earliest-printed first, so multiple brand-new PDOs destined for the
        # same station land in the queue in the correct FIFO print-time order
        # within this single pass.
        candidates = sorted(prints.items(), key=lambda kv: kv[1][1])
        changed = False
        # Guards 1 & 2 (unfinished, not-already-on-board) fire for most of the
        # 30-day print window every poll, so they stay SILENT — logging them just
        # floods the console. Only the guards that block an order a GL would expect
        # to SEE (suppressed / station-not-configured) are logged below.
        for po, (sql_station, print_ts) in candidates:
            if pdo_by_id(po) is None:
                continue  # not active/unfinished on the board (remaining ≤ 0, or unknown PO) — skip
            station, _bin = find_pdo_current_location(po, assignments)
            if station is not None:
                continue  # already assigned somewhere — not "brand new" anymore
            supp_iso = suppress.get(po)
            if supp_iso:
                try:
                    if print_ts <= datetime.fromisoformat(supp_iso):
                        print(f'[auto-queue skip] {po}: manually removed — suppressed until a print newer than {supp_iso}')
                        continue  # GL removed this; wait for a NEWER LP print before re-adding
                except Exception:
                    pass
            board_station = normalize_station_board(sql_station)
            if board_station not in station_ids:
                print(f'[auto-queue skip] {po}: printed station {sql_station!r} → {board_station!r} is not a configured station on this board ({sorted(station_ids)})')
                continue  # printed station isn't a real/configured station on this board
            bins = assignments.get(board_station)
            if bins is None:
                print(f'[auto-queue skip] {po}: station {board_station} has no assignment bin')
                continue
            if not bins.get('inprogress'):
                bins['inprogress'] = po
                print(f'[auto-queue] {board_station}: {po} printed {print_ts} → placed in-progress (station was idle)')
            else:
                bins.setdefault('queue', []).append(po)
                print(f'[auto-queue] {board_station}: {po} printed {print_ts} → appended to queue')
            changed = True
        if changed:
            save_assignments(assignments)
            rebuild_final_pdo_cache()

def sync_timer_pause_state(station, worker_count, timers):
    """Keep a station's timer 'paused' flag in sync with whether it currently
    has any worker staffed.
    - 0 workers → paused=True. The ETC display goes to 'Paused — no worker'
      instead of continuing to tick down as if someone were still working.
    - Workers present again after being paused → RESUME: reset started_at to
      NOW and restart the fg_count. This is the important part — without it,
      the clock would resume counting from whenever the order was originally
      assigned, which could be hours or days stale (overnight, shift gap),
      making it show 'OVERDUE' the instant someone gets reassigned even
      though realistically work just started again.
    No-op if the station has no active timer (nothing in-progress there).
    Returns True if it actually changed the timer (so callers can avoid a
    pointless disk write when nothing moved)."""
    timer = timers.get(station)
    if timer is None:
        return False
    if worker_count == 0:
        if not timer.get('paused'):
            timer['paused'] = True
            return True
        return False
    else:
        if timer.get('paused'):
            timer['started_at'] = datetime.now().isoformat()
            timer['fg_count'] = 0
            timer['paused'] = False
            return True
        return False

def reset_all_timers_for_new_shift():
    """Reset the ETC timer on EVERY station with an in-progress order — set
    started_at to now and clear fg_count — so ETC is recalculated fresh from
    each PDO's current remaining_qty for the new shift. Called at each shift
    boundary (Ca1 end / morning start). Whether a station's clock actually runs
    afterwards still depends on active workers + production hours (paused flag)."""
    with _cache_lock:
        assignments  = load_assignments()
        timers       = load_timers()
        workers_asgn = load_workers_assigned()
        now_iso = datetime.now().isoformat()
        changed = False
        for station, bins in assignments.items():
            if bins.get('inprogress'):
                timers[station] = {
                    'started_at': now_iso,
                    'fg_count': 0,
                    'paused': _count_active_workers(station, workers_asgn) == 0,
                }
                changed = True
        if changed:
            save_timers(timers)
        print('[shift reset] ETC recalculated from remaining qty for all in-progress stations')

def reset_daily_worker_staffing():
    """Clear carried-over worker→station staffing at the start of each production
    day. Worker assignments are per-shift: normally prune_expired_workers empties a
    station's operators when their shift ends, so stations sit empty of operators
    between shifts and the GL re-staffs every morning. But that self-clean only
    happens if the board is actually running/pruning across the shift-end boundary.
    If it wasn't — shut down overnight, or the attendance read was failing at
    shift end (e.g. the all-HC day: HC ends 16:30, and if nothing pruned then) —
    yesterday's operators are never cleared, and this morning the same people, now
    on a different shift (yesterday HC → today Ca1) and 'available now', get KEPT
    by the prune on yesterday's stations. A once-per-day clear closes that gap.

    It only removes stale worker staffing (never PDO assignments), and is a no-op
    on any day the shift-end prune already did its job. Persisted by date so it
    fires once per calendar day even across a restart."""
    today = date.today().isoformat()
    state = load_json(f'worker_reset_{AREA}.json', {})
    if state.get('date') == today:
        return
    with _cache_lock:
        workers_asgn = load_workers_assigned()
        if any(workers_asgn.values()):
            save_workers_assigned({sid: [] for sid in workers_asgn})
            print(f'[daily reset] new production day {today} — cleared carried-over worker '
                  f'staffing; operators re-staffed per today\'s shift (PDOs untouched)')
    save_json(f'worker_reset_{AREA}.json', {'date': today})

def prune_expired_workers():
    """Remove workers from station assignments once their shift window ends —
    they also naturally disappear from the available pool via get_available_workers.
    Also pauses the ETC timer for any station that just lost its last worker
    this way (shift end), so the clock properly stops rather than continuing
    to tick as if production were still happening.

    GUARD: only prune when TODAY's attendance actually read (_last_attendance_ok).
    A failed/missing/unsynced attendance read returns an empty roster that is
    indistinguishable from 'nobody on shift' — so without this, a fresh restart
    while the attendance file is briefly unreadable strips EVERY assigned operator
    off every station and saves it. Skip and retry next poll once it's readable."""
    if not _last_attendance_ok:
        print('[shift-end] skipped — today\'s attendance could not be read; '
              'refusing to remove operators against an empty/failed roster')
        return
    # A worker stays assigned only while they're in the live-available window —
    # the SAME check the board pool uses (_worker_available_now), so the two can
    # never disagree. (The old `now <= end` check had no lower bound and broke
    # across midnight: after a Ca2 shift ended at 22:20, the next morning's small
    # `now` was < the 22.333 end, so the worker was wrongly kept and showed as a
    # bare ID on the station.)
    _now_dec = datetime.now().hour + datetime.now().minute / 60
    available_ids = {str(w['id']) for w in _worker_cache
                     if _worker_available_now(w, _now_dec)}
    active_ids = _active_worker_id_set()
    workers_asgn = load_workers_assigned()
    timers = load_timers()
    changed = False
    timers_changed = False
    for station, ids in workers_asgn.items():
        new_ids = [i for i in ids if str(i) in available_ids]
        if len(new_ids) != len(ids):
            removed = set(ids) - set(new_ids)
            print(f'[shift-end] {station}: removing expired workers {removed}')
            workers_asgn[station] = new_ids
            changed = True
        # Pause/resume off the ACTIVE count so the clock starts the moment a
        # pre-assigned upcoming worker's shift actually begins (started_at resets
        # to now via sync_timer_pause_state's resume branch).
        active_count = sum(1 for wid in new_ids if str(wid) in active_ids)
        if sync_timer_pause_state(station, active_count, timers):
            timers_changed = True
    if changed:
        save_workers_assigned(workers_asgn)
    if timers_changed:
        save_timers(timers)

# ── Write-back to Excel ────────────────────────────────────────────────────────
def write_station_to_excel(pdo_id, station_id):
    """DISABLED — openpyxl's keep_vba=True save can corrupt complex .xlsm files
    (custom UI, form controls, certain macro structures) on repeated writes.
    The JSON files in /data are already the source of truth for the board;
    this was only a nice-to-have sync back to Excel, never required.
    Re-enable only after finding a write method proven safe for this specific file."""
    pass

# ── Routes ─────────────────────────────────────────────────────────────────────
@app.route('/')
def index():
    return render_template('board.html', area=AREA)

@app.route('/board/<area>')
def board(area):
    return render_template('board.html', area=area)

# ── API: full board state ──────────────────────────────────────────────────────
def build_board_snapshot(area):
    """Builds the full board state dict — same shape /api/board/<area> returns.
    Factored out so both the HTTP route AND the shared-file dashboard writer
    (which needs this same data but isn't an HTTP request) can use it."""
    with _cache_lock:
        assignments  = load_assignments()
        workers_asgn = load_workers_assigned()
        timers       = load_timers()

        # NOTE: snapshot building is now READ-ONLY — it no longer prunes/saves.
        # Pruning used to run here, on every board GET (~10s) and every dashboard
        # write (~5s); any brief cache hiccup during one of those frequent calls
        # would clear a station's in-progress order and persist it, which showed
        # up as orders "randomly" bouncing back to the queue. Pruning is now done
        # ONLY in the 60s poll (refresh_fg_and_advance), right after the cache is
        # fully rebuilt from fresh Excel+SQL, so it acts on a consistent view.
        # A stale id left in assignments here is harmless for display: a missing
        # in-progress renders as an empty bin, and missing queue ids are filtered
        # out below — the next poll cleans them up authoritatively.

        assigned_ids = set()
        for st in assignments.values():
            if st['inprogress']: assigned_ids.add(st['inprogress'])
            for qid in st.get('queue', []):
                assigned_ids.add(qid)

        queue_pool = [p for p in _pdo_cache if p['id'] not in assigned_ids]

        # Live pool: filter the roster by the current time (15-min pre-open →
        # shift end) and stamp the live 'active' flag (upcoming next-shift workers
        # show active=False so the UI marks them SOON and ETC ignores them). Both
        # membership and active are recomputed here every board GET, so they track
        # the clock to the minute regardless of the ~5-min attendance refresh.
        _act = _active_worker_id_set()
        _now_dec = datetime.now().hour + datetime.now().minute / 60
        _avail_ids = {str(w['id']) for w in _worker_cache if _worker_available_now(w, _now_dec)}
        workers_out = [{**w, 'active': str(w['id']) in _act}
                       for w in _worker_cache if str(w['id']) in _avail_ids]

        in_prod = is_production_time()  # ETC freezes outside production hours

        board_out = {}
        for sid in _station_cache:
            bins      = assignments.get(sid, {'inprogress': None, 'queue': []})
            ip_pdo    = pdo_by_id(bins['inprogress']) if bins['inprogress'] else None
            queue_ids = bins.get('queue', [])
            queue_pdos = [pdo_by_id(qid) for qid in queue_ids]
            queue_pdos = [p for p in queue_pdos if p]  # drop any not found
            # Only show workers who are currently available (in the live pool);
            # an operator whose shift has ended vanishes from the station display
            # immediately, not just after the next 60s prune. (Their names resolve
            # from the pool, so a lingering assignment would otherwise show a bare
            # ID.) prune_expired_workers clears them from the saved assignment.
            workers   = [wid for wid in workers_asgn.get(sid, []) if str(wid) in _avail_ids]
            timer     = timers.get(sid, {})
            # ETC counts only workers whose shift has actually started — a
            # pre-assigned upcoming (next-shift) worker shouldn't make the clock
            # tick until they're really on the floor.
            active_count = _count_active_workers(sid, workers_asgn)
            # The ETC clock only RUNS during production hours with an active
            # worker. Outside production hours it FREEZES (off_hours); with no
            # active worker it's paused. Either way the countdown stops ticking.
            running = bool(ip_pdo) and active_count > 0 and in_prod

            etc_iso = None
            if running and timer.get('started_at'):
                etc_iso = calc_etc(
                    ip_pdo['finished_hrs'],
                    ip_pdo['remaining_qty'],
                    ip_pdo['qty'],
                    active_count,
                    timer['started_at']
                )

            clear_iso = None
            if running and (ip_pdo or queue_pdos):
                clear_iso = calc_queue_clear(
                    ip_pdo, queue_pdos, active_count,
                    timer.get('started_at') or datetime.now().isoformat()
                )

            # Static remaining-work estimate shown while frozen/paused (so the
            # value is held, not blanked). Uses active workers, or the assigned
            # count as a fallback so an unstaffed order still shows an estimate.
            remaining_min = None
            if ip_pdo and ip_pdo.get('qty') and ip_pdo.get('finished_hrs'):
                wc = active_count if active_count > 0 else len(workers)
                if wc > 0:
                    hrs_per_unit = ip_pdo['finished_hrs'] / ip_pdo['qty']
                    remaining_min = round(hrs_per_unit * ip_pdo.get('remaining_qty', 0) / wc * 60)

            off_hours = bool(ip_pdo) and not in_prod
            # Paused: in progress but the clock isn't running (no active worker,
            # or outside production hours).
            paused = bool(ip_pdo) and not running

            board_out[sid] = {
                'inprogress': ip_pdo,
                'queue':      queue_pdos,
                'workers':    workers,
                'timer': {
                    'started_at':   timer.get('started_at'),
                    'etc':          etc_iso,
                    'queue_clear':  clear_iso,
                    'fg_count':     timer.get('fg_count', 0),
                    'paused':       paused,
                    'off_hours':    off_hours,
                    'remaining_min': remaining_min,
                }
            }

        return {
            'area':          area,
            'area_label':    _area_label,
            'stations':      _station_cache,
            'workers':       workers_out,
            'queue':         queue_pool,
            'board':         board_out,
            'pdo_status':    _last_pdo_status,
            'fg_status':     _last_fg_status,
            'ts_status':     _last_ts_status,
            'shift_windows': _shift_windows_cache,
            'capacity_hours': _day_capacity_hours,
            'in_production': in_prod,
            'dev_mode':      DEV_MODE,
            'ts':            datetime.now().strftime('%H:%M:%S'),
        }

@app.route('/api/board/<area>')
def api_board(area):
    if area not in AREAS:
        return jsonify({'error': 'unknown area'}), 404
    threading.Thread(target=refresh_cache, daemon=True).start()
    return jsonify(build_board_snapshot(area))

# ── API: report view — full PDO list with live assignment status ───────────────
@app.route('/api/report/pdos')
def api_report_pdos():
    try:
        with _cache_lock:
            assignments = load_assignments()
            result = []
            for p in _pdo_cache:
                station, bin_type = find_pdo_current_location(p['id'], assignments)
                p2 = dict(p)
                p2['assigned_station'] = station
                p2['assigned_bin'] = bin_type  # 'inprogress' | 'queue' | None
                # Green tag when this PDO has a License Plate print record.
                p2['has_lp'] = (p['id'] in _license_plate_cache) or bool(p.get('parent_id') and p['parent_id'] in _license_plate_cache)
                result.append(p2)
        return jsonify({'ok': True, 'pdos': result, 'stations': _station_cache,
                        'capacity_hours': _day_capacity_hours})
    except Exception as e:
        print(f'[api_report_pdos error] {e}')
        return jsonify({'error': str(e)}), 500

_capday_cache = {}   # (from_iso, to_iso) → (map, ts), TTL = worker-refresh throttle

@app.route('/api/report/capacity_by_day')
def api_capacity_by_day():
    """Per-day available capacity for the report's hours-chart reference line.
    Returns {YYYY-MM-DD: hours} for [from, to] (defaults to today-14 … today+30),
    omitting days with no attendance filled so the chart leaves a gap there.
    Cached for the worker-refresh interval; reads the attendance file outside the
    cache lock."""
    frm = request.args.get('from')
    to  = request.args.get('to')
    try:
        d_from = datetime.strptime(frm, '%Y-%m-%d').date() if frm else date.today() - timedelta(days=14)
        d_to   = datetime.strptime(to,  '%Y-%m-%d').date() if to  else date.today() + timedelta(days=30)
    except ValueError:
        return jsonify({'error': 'bad date (expected YYYY-MM-DD)'}), 400
    if d_to < d_from:
        return jsonify({'ok': True, 'capacity_by_day': {}})
    if (d_to - d_from).days > 120:      # bound the scan so a huge range can't stall it
        d_to = d_from + timedelta(days=120)
    key = (d_from.isoformat(), d_to.isoformat())
    now = time.time()
    cached = _capday_cache.get(key)
    if cached and (now - cached[1]) < WORKER_REFRESH_THROTTLE_SEC:
        return jsonify({'ok': True, 'capacity_by_day': cached[0]})
    try:
        m = compute_capacity_by_day(AREA, d_from, d_to)
    except Exception as e:
        print(f'[api_capacity_by_day error] {e}')
        return jsonify({'error': str(e)}), 500
    _capday_cache[key] = (m, now)
    return jsonify({'ok': True, 'capacity_by_day': m})

# ── API: report view — assign / move / unassign a PDO with smart placement ─────
@app.route('/api/report/set_station', methods=['POST'])
def api_report_set_station():
    """Single endpoint handling assign, edit (move), and remove (unassign) —
    matches the report view's one-dropdown-does-everything UX.
    Placement is 'smart': the target station's in-progress slot if it's
    empty, otherwise appended to that station's queue — same as dragging a
    card onto a station on the main board."""
    d = request.get_json(silent=True) or {}
    pdo_id = d.get('pdo_id')
    new_station = d.get('station') or None  # None/blank = unassign
    if not pdo_id:
        return jsonify({'error': 'missing pdo_id'}), 400
    if new_station and new_station not in _station_cache:
        return jsonify({'error': f'unknown station {new_station}'}), 400

    try:
        with _cache_lock:
            pdo = pdo_by_id(pdo_id)
            if not pdo:
                return jsonify({'error': f'{pdo_id} not found or already complete'}), 404

            assignments = load_assignments()
            timers      = load_timers()
            workers_asgn = load_workers_assigned()

            cur_station, cur_bin = find_pdo_current_location(pdo_id, assignments)

            # Remove from wherever it currently sits (no-op if unassigned)
            if cur_station:
                if cur_bin == 'inprogress':
                    assignments[cur_station]['inprogress'] = None
                    timers.pop(cur_station, None)
                else:
                    q = assignments[cur_station].get('queue', [])
                    if pdo_id in q:
                        q.remove(pdo_id)

            if new_station:
                if new_station not in assignments:
                    assignments[new_station] = {'inprogress': None, 'queue': []}
                if assignments[new_station]['inprogress'] is None:
                    # Empty slot — goes straight to in-progress
                    assignments[new_station]['inprogress'] = pdo_id
                    timers[new_station] = {'started_at': datetime.now().isoformat(), 'fg_count': 0}
                    if _count_active_workers(new_station, workers_asgn) == 0:
                        timers[new_station]['paused'] = True
                else:
                    # Occupied — goes to the back of that station's queue
                    assignments[new_station].setdefault('queue', []).append(pdo_id)
                update_split_station(pdo_id, new_station)
            elif cur_station:
                # Manual unassign (dropdown → blank): suppress LP auto-requeue
                # until a newer License Plate print arrives for this PDO.
                _record_lp_suppression(pdo_id)

            save_assignments(assignments)
            save_timers(timers)
            write_station_to_excel(pdo_id, new_station or '')

        return jsonify({'ok': True, 'station': new_station,
                         'placement': 'inprogress' if (new_station and assignments[new_station]['inprogress'] == pdo_id) else ('queue' if new_station else None)})
    except Exception as e:
        print(f'[api_report_set_station error] {e}')
        import traceback; traceback.print_exc()
        return jsonify({'error': str(e)}), 500

# ── API: Finished PDOs report (Excel data + SQL finish date) ───────────────────
def get_fg_finish_dates(po_numbers):
    """{po: last_scan_datetime} — MAX([Input time]) per PO across ALL time (not
    windowed), i.e. when the LAST box of that PO was scanned. Used only for the
    finished-PDOs report's Finish Date column, for the (small) set of finished
    POs, so the unbounded MAX is cheap."""
    if not PYODBC_AVAILABLE or not po_numbers:
        return {}
    try:
        conn = pyodbc.connect(SQL_CONN_PROD, timeout=15)
        cursor = conn.cursor()
        result = {}
        pos = [str(p).strip() for p in po_numbers if p]
        for i in range(0, len(pos), 1000):
            chunk = pos[i:i+1000]
            ph = ','.join('?' for _ in chunk)
            cursor.execute(
                f"SELECT PO, MAX([Input time]) FROM [dbo].[FG_Database_All] "
                f"WHERE PO IN ({ph}) GROUP BY PO", *chunk)
            for r in cursor.fetchall():
                if r[1] is not None:
                    result[str(r[0]).strip()] = r[1]
        conn.close()
        return result
    except Exception as e:
        print(f'[FG finish-date SQL error] {e}')
        return {}

@app.route('/api/report/finished')
def api_report_finished():
    """All FULLY-FINISHED PDOs (SQL Finished Qty >= Excel Order Qty). Row data
    comes from the Excel plan (A/B/D/E/F/G/H/I/L/N + the K 'sample' flag) and the
    Finish Date is when the last box was scanned in FG_Database_All."""
    try:
        # Snapshot the finished orders' Excel fields under the lock, then do the
        # SQL finish-date lookup outside it (never hold the cache lock during SQL).
        with _cache_lock:
            snap = []
            for p in _raw_pdo_cache:
                qty = p.get('qty') or 0
                if qty <= 0:
                    continue
                completed = _fg_completed_cache.get(p['id'], 0)
                if completed >= qty:  # Finished Qty (SQL) == Order Qty (fully done)
                    snap.append({
                        'id': p['id'], 'part': p.get('part', ''), 'desc': p.get('desc', ''),
                        'qty': qty, 'dest': p.get('dest', ''), 'pack_type': p.get('pack_type', ''),
                        'receive_date': p.get('receive_date', ''), 'ship_date': p.get('ship_date', ''),
                        'note': p.get('note', ''), 'vendor_name': p.get('vendor_name', ''),
                        'is_sample': p.get('is_sample', False),
                    })
        finish_dates = get_fg_finish_dates([r['id'] for r in snap])
        for r in snap:
            fd = finish_dates.get(r['id'])
            r['finish_date'] = str(fd) if fd else ''
        return jsonify({'ok': True, 'rows': snap})
    except Exception as e:
        print(f'[api_report_finished error] {e}')
        return jsonify({'error': str(e)}), 500

# ── API: GL manual priority override for a PDO ──────────────────────────────────
@app.route('/api/priority', methods=['POST'])
def api_priority():
    """Set/clear a manual priority override (ph/pm/pl). Empty/None clears it and
    reverts the order to its computed priority."""
    d = request.get_json(silent=True) or {}
    pdo_id = d.get('pdo_id')
    pr = d.get('priority')
    if not pdo_id:
        return jsonify({'error': 'missing pdo_id'}), 400
    if pr not in ('ph', 'pm', 'pl', None, ''):
        return jsonify({'error': 'invalid priority'}), 400
    try:
        with _cache_lock:
            overrides = load_json(f'priority_{AREA}.json', {})
            if pr in ('ph', 'pm', 'pl'):
                overrides[pdo_id] = pr
            else:
                overrides.pop(pdo_id, None)
            save_json(f'priority_{AREA}.json', overrides)
            rebuild_final_pdo_cache()
        return jsonify({'ok': True, 'priority': pr or None})
    except Exception as e:
        print(f'[api_priority error] {e}')
        return jsonify({'error': str(e)}), 500

# ── API: refresh PDO list from Excel ──────────────────────────────────────────
# ── API: full-day attendance roster for the Gantt view ──────────────────────────
@app.route('/api/attendance/gantt')
def api_attendance_gantt():
    try:
        # include_lent_away=True so fully-lent operators still appear on the
        # report (flagged, with a note) even though they're kept out of the
        # assignable pool and the capacity figure.
        info = read_attendance_shift_info(AREA, date.today(), include_lent_away=True)
        rows = []
        for w in info:
            start, end = w.get('start'), w.get('end')
            if start is None or end is None:
                start, end = _worker_window(w['shift'], w['ot_clock'])
            lend_hours = w.get('lend_hours', 0.0)
            note_kind = w.get('note_kind', '')
            if note_kind == 'partial':
                borrow_note = 'note_window'    # explicit time-window note for this area
            elif note_kind == 'away' or w.get('lent_away'):
                borrow_note = 'lent_away'      # full shift → in another area all day
            elif w.get('is_borrow'):
                borrow_note = 'borrowed'       # borrowed IN from another area
            elif lend_hours:
                borrow_note = 'lent_partial'   # here, but lent out for part of the day
            else:
                borrow_note = ''
            rows.append({
                'id':          w['id'],
                'name':        w['name'],
                'shift':       w['shift'],
                'ot_clock':    w['ot_clock'],
                'start':       start,
                'end':         end,
                'lent_away':   bool(w.get('lent_away')),
                'is_borrow':   bool(w.get('is_borrow')),
                'lend_hours':  round(abs(lend_hours), 2) if lend_hours else 0,
                'cap_hours':   w.get('cap_hours', 0),
                'windows':     w.get('windows', []),
                'note':        w.get('note', ''),
                'borrow_note': borrow_note,
            })
        # Sort by shift group (Ca1, Ca2, HC), then by start time within each
        shift_order = {'Ca 1': 0, 'Ca 2': 1, 'HC': 2}
        rows.sort(key=lambda r: (shift_order.get(r['shift'], 9), r['start'], r['name']))
        return jsonify({'ok': True, 'rows': rows, 'date': date.today().strftime('%Y-%m-%d')})
    except Exception as e:
        print(f'[api_attendance_gantt error] {e}')
        return jsonify({'error': str(e)}), 500

# ── API: station management (add/remove) ───────────────────────────────────────
@app.route('/api/stations/add', methods=['POST'])
def api_station_add():
    d = request.get_json(silent=True) or {}
    sid = (d.get('station_id') or '').strip()
    if not sid:
        return jsonify({'error': 'station id required'}), 400
    with _cache_lock:
        stations = load_stations()
        if sid in stations:
            return jsonify({'error': f'{sid} already exists'}), 409
        stations.append(sid)
        save_stations(stations)
        global _station_cache
        _station_cache = stations
        # Initialize an empty assignment slot so the board has something to render
        assignments = load_assignments()
        assignments[sid] = {'inprogress': None, 'queue': []}
        save_assignments(assignments)
    return jsonify({'ok': True, 'stations': stations})

@app.route('/api/stations/remove', methods=['POST'])
def api_station_remove():
    d = request.get_json(silent=True) or {}
    sid = (d.get('station_id') or '').strip()
    force = bool(d.get('force', False))
    if not sid:
        return jsonify({'error': 'station id required'}), 400
    with _cache_lock:
        stations = load_stations()
        if sid not in stations:
            return jsonify({'error': f'{sid} not found'}), 404

        assignments  = load_assignments()
        bins         = assignments.get(sid, {})
        workers_asgn = load_workers_assigned()
        has_order    = bool(bins.get('inprogress')) or bool(bins.get('queue'))
        has_workers  = bool(workers_asgn.get(sid))

        if (has_order or has_workers) and not force:
            reason = []
            if has_order: reason.append('an assigned/queued order')
            if has_workers: reason.append('assigned worker(s)')
            return jsonify({
                'error': f'{sid} still has {" and ".join(reason)} — clear it first, or resend with force=true',
                'has_order': has_order,
                'has_workers': has_workers,
            }), 409

        stations.remove(sid)
        save_stations(stations)
        assignments.pop(sid, None)
        save_assignments(assignments)
        workers_asgn.pop(sid, None)
        save_workers_assigned(workers_asgn)
        timers = load_timers()
        timers.pop(sid, None)
        save_timers(timers)
        global _station_cache
        _station_cache = stations
    return jsonify({'ok': True, 'stations': stations})

# ── API: refresh PDO list from Excel ────────────────────────────────────────────
@app.route('/api/refresh_pdos', methods=['POST'])
def api_refresh():
    # Manual "Refresh orders" is a FULL gather, not just an Excel re-read.
    # Force the Excel read, then run the exact same FG/WIP/scan pull +
    # License Plate auto-queue + auto-advance that the 60s poll does — all
    # synchronously, before we return. That way anything the License Plate
    # query auto-assigns to a station is already in _pdo_cache/assignments by
    # the time the board re-renders, instead of only appearing on the next
    # 60s poll a few seconds later.
    refresh_cache(force=True)
    try:
        refresh_fg_and_advance()
    except Exception as e:
        print(f'[api_refresh FG/License-Plate pass error] {e}')
    return jsonify({
        'ok': True,
        'count': len(_pdo_cache),
        'status': _last_pdo_status,
        'fg_status': _last_fg_status,
        'ts': datetime.now().strftime('%H:%M:%S')
    })

@app.route('/api/refresh_fg', methods=['POST'])
def api_refresh_fg():
    """On-demand FG pull — normally happens automatically every 60s,
    this lets GL force it immediately (e.g. right after a known scan)."""
    refresh_fg_and_advance()
    return jsonify({
        'ok': True,
        'fg_status': _last_fg_status,
        'count': len(_pdo_cache),
        'ts': datetime.now().strftime('%H:%M:%S')
    })

# ── API: assign PDO to bin ────────────────────────────────────────────────────
@app.route('/api/assign', methods=['POST'])
def api_assign():
    d = request.get_json(silent=True) or {}
    station, bin_type, pdo_id = d.get('station'), d.get('bin'), d.get('pdo_id')
    if not all([station, bin_type, pdo_id]):
        return jsonify({'error': 'missing fields'}), 400
    if bin_type not in ('inprogress', 'queue'):
        return jsonify({'error': 'invalid bin type'}), 400

    try:
        with _cache_lock:
            refresh_cache()
            assignments = load_assignments()
            timers      = load_timers()
            prune_stale_assignments(assignments, timers)

            if station not in assignments:
                assignments[station] = {'inprogress': None, 'queue': []}

            if bin_type == 'inprogress' and assignments[station]['inprogress']:
                return jsonify({'error': f'in-progress bin occupied on {station}'}), 409

            # Remove from any current slot across all stations (inprogress or queue)
            prev_station = None
            for st, bins in assignments.items():
                if bins['inprogress'] == pdo_id:
                    bins['inprogress'] = None
                    prev_station = st
                if pdo_id in bins.get('queue', []):
                    bins['queue'].remove(pdo_id)

            if prev_station and prev_station in timers:
                timers.pop(prev_station)

            if bin_type == 'inprogress':
                assignments[station]['inprogress'] = pdo_id
                timers[station] = {
                    'started_at': datetime.now().isoformat(),
                    'fg_count':   0,
                }
                # If nobody's staffing this station yet, start the timer paused —
                # otherwise it would show a ticking ETC before anyone's actually working it.
                workers_asgn = load_workers_assigned()
                worker_count = _count_active_workers(station, workers_asgn)
                if worker_count == 0:
                    timers[station]['paused'] = True
            else:  # queue — append to end (drop order = queue order)
                assignments[station].setdefault('queue', []).append(pdo_id)

            save_assignments(assignments)
            save_timers(timers)
            write_station_to_excel(pdo_id, station)
            update_split_station(pdo_id, station)
        return jsonify({'ok': True})

    except Exception as e:
        print(f'[api_assign error] {e}')
        import traceback; traceback.print_exc()
        return jsonify({'error': str(e)}), 500

# ── API: split a PDO across up to 4 stations ────────────────────────────────────
@app.route('/api/splits/create', methods=['POST'])
def api_splits_create():
    d = request.get_json(silent=True) or {}
    parent_id = d.get('parent_id')
    splits_input = d.get('splits')  # [{letter, station (nullable), qty}, ...]
    if not parent_id or not splits_input:
        return jsonify({'error': 'missing fields'}), 400
    if not (2 <= len(splits_input) <= 4):
        return jsonify({'error': 'must specify 2 to 4 splits'}), 400

    letters = [s.get('letter') for s in splits_input]
    if len(set(letters)) != len(letters) or not all(l in ('A', 'B', 'C', 'D') for l in letters):
        return jsonify({'error': 'split letters must be unique, from A-D'}), 400

    # Normalize blank/empty station values to None for consistent handling
    for s in splits_input:
        if not s.get('station'):
            s['station'] = None

    stations_nonblank = [s['station'] for s in splits_input if s['station']]
    if len(set(stations_nonblank)) != len(stations_nonblank):
        return jsonify({'error': 'each split must go to a different station'}), 400
    for st in stations_nonblank:
        if st not in _station_cache:
            return jsonify({'error': f'unknown station {st}'}), 400

    try:
        with _cache_lock:
            pdo = pdo_by_id(parent_id)
            if not pdo:
                return jsonify({'error': f'{parent_id} not found or already complete'}), 404
            if pdo.get('is_split_child'):
                return jsonify({'error': 'cannot split a split — merge it back first'}), 400

            existing_splits = load_splits()
            if parent_id in existing_splits:
                return jsonify({'error': f'{parent_id} already has active splits — merge first'}), 409

            remaining_qty = pdo['remaining_qty']
            try:
                qtys = [int(s.get('qty', 0)) for s in splits_input]
            except (TypeError, ValueError):
                return jsonify({'error': 'invalid quantity'}), 400
            if any(q <= 0 for q in qtys):
                return jsonify({'error': 'every split needs a quantity greater than 0'}), 400
            entered_total = sum(qtys)
            if entered_total != remaining_qty:
                return jsonify({'error': f'split total ({entered_total}) must exactly equal remaining qty ({remaining_qty})'}), 400

            assignments  = load_assignments()
            timers       = load_timers()
            workers_asgn = load_workers_assigned()

            cur_station, cur_bin = find_pdo_current_location(parent_id, assignments)

            # Station requirement depends on whether the order is currently
            # placed anywhere:
            #  - Currently assigned (cur_station set): split 'A' auto-inherits
            #    that station if left blank; every OTHER split (B/C/D) MUST
            #    have an explicit station — leaving work half-assigned while
            #    part of it is already physically running doesn't make sense.
            #  - Not currently assigned: every split may be left blank, in
            #    which case it just sits in the general pool unplaced, for GL
            #    to drag onto a station later.
            if cur_station:
                for s in splits_input:
                    if s['letter'] == 'A' and not s['station']:
                        s['station'] = cur_station
                for s in splits_input:
                    if not s['station']:
                        return jsonify({'error': f'split {s["letter"]} needs a station — {parent_id} is currently in progress, so only split A can auto-continue there'}), 400

            parent_completed_global = pdo['completed_qty']
            splits_record = {'splits': {}}

            for s in splits_input:
                letter      = s['letter']
                station     = s['station']  # may still be None here
                entered_qty = int(s['qty'])
                is_inheriting = bool(station) and (station == cur_station)
                credit    = parent_completed_global if is_inheriting else 0
                total_qty = credit + entered_qty
                now_iso = datetime.now().isoformat()
                splits_record['splits'][letter] = {
                    'station': station,
                    'station_set_at': now_iso if station else None,
                    'credit': credit,
                    'total_qty': total_qty,
                }

                if not station:
                    continue  # left blank — stays unplaced in the general pool

                split_id = f"{parent_id}-{letter}"

                if is_inheriting and cur_bin:
                    # Same station already had this order — repoint the existing
                    # slot in place, the physical work just continues uninterrupted.
                    if cur_bin == 'inprogress':
                        assignments[cur_station]['inprogress'] = split_id
                        timers[cur_station] = {'started_at': datetime.now().isoformat(), 'fg_count': 0}
                        if _count_active_workers(cur_station, workers_asgn) == 0:
                            timers[cur_station]['paused'] = True
                    else:
                        q = assignments[cur_station]['queue']
                        if parent_id in q:
                            q[q.index(parent_id)] = split_id
                        else:
                            q.append(split_id)
                else:
                    # Fresh assignment to the chosen station, normal placement rules
                    if station not in assignments:
                        assignments[station] = {'inprogress': None, 'queue': []}
                    if assignments[station]['inprogress'] is None:
                        assignments[station]['inprogress'] = split_id
                        timers[station] = {'started_at': datetime.now().isoformat(), 'fg_count': 0}
                        if _count_active_workers(station, workers_asgn) == 0:
                            timers[station]['paused'] = True
                    else:
                        assignments[station].setdefault('queue', []).append(split_id)

            # Safety-net cleanup — clear parent_id from anywhere it might still
            # be referenced (the in-place repoint above already handles the
            # common case, this just guards against anything left over).
            for st, bins in assignments.items():
                if bins.get('inprogress') == parent_id:
                    bins['inprogress'] = None
                    timers.pop(st, None)
                if parent_id in bins.get('queue', []):
                    bins['queue'].remove(parent_id)

            save_assignments(assignments)
            save_timers(timers)
            existing_splits[parent_id] = splits_record
            save_splits(existing_splits)

            refresh_split_completions()
            rebuild_final_pdo_cache()

        return jsonify({'ok': True, 'splits': splits_record})
    except Exception as e:
        print(f'[api_splits_create error] {e}')
        import traceback; traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/splits/merge', methods=['POST'])
def api_splits_merge():
    d = request.get_json(silent=True) or {}
    parent_id = d.get('parent_id')
    target_station = d.get('target_station')  # optional — None/blank = release to pool
    if not parent_id:
        return jsonify({'error': 'missing parent_id'}), 400
    if target_station and target_station not in _station_cache:
        return jsonify({'error': f'unknown station {target_station}'}), 400

    try:
        with _cache_lock:
            existing_splits = load_splits()
            if parent_id not in existing_splits:
                return jsonify({'error': f'{parent_id} has no active splits'}), 404

            assignments  = load_assignments()
            timers       = load_timers()
            workers_asgn = load_workers_assigned()
            info = existing_splits[parent_id]

            cleared_stations = set()
            for letter in info.get('splits', {}):
                split_id = f"{parent_id}-{letter}"
                for st, bins in assignments.items():
                    if bins.get('inprogress') == split_id:
                        bins['inprogress'] = None
                        timers.pop(st, None)
                        cleared_stations.add(st)
                    if split_id in bins.get('queue', []):
                        bins['queue'].remove(split_id)

            if target_station:
                if target_station not in assignments:
                    assignments[target_station] = {'inprogress': None, 'queue': []}
                if assignments[target_station]['inprogress'] is None:
                    assignments[target_station]['inprogress'] = parent_id
                    timers[target_station] = {'started_at': datetime.now().isoformat(), 'fg_count': 0}
                    if _count_active_workers(target_station, workers_asgn) == 0:
                        timers[target_station]['paused'] = True
                else:
                    assignments[target_station].setdefault('queue', []).append(parent_id)

            # A split that was IN-PROGRESS on a station left that slot empty when
            # merged. If the merged parent didn't take that slot, promote the
            # station's next queued order — the same auto-advance a manual removal
            # does. Without this the next PDO never advanced after a merge.
            for st in cleared_stations:
                if assignments[st].get('inprogress') is None and assignments[st].get('queue'):
                    nxt = assignments[st]['queue'].pop(0)
                    assignments[st]['inprogress'] = nxt
                    timers[st] = {'started_at': datetime.now().isoformat(), 'fg_count': 0}
                    if _count_active_workers(st, workers_asgn) == 0:
                        timers[st]['paused'] = True

            save_assignments(assignments)
            save_timers(timers)

            del existing_splits[parent_id]
            save_splits(existing_splits)

            refresh_split_completions()
            rebuild_final_pdo_cache()

        return jsonify({'ok': True})
    except Exception as e:
        print(f'[api_splits_merge error] {e}')
        import traceback; traceback.print_exc()
        return jsonify({'error': str(e)}), 500

# ── API: move PDO between bins ────────────────────────────────────────────────
@app.route('/api/move', methods=['POST'])
def api_move():
    d = request.get_json(silent=True) or {}
    from_s, from_b = d.get('from_station'), d.get('from_bin')
    to_s,   to_b   = d.get('to_station'),   d.get('to_bin')
    pdo_id_hint    = d.get('pdo_id')          # optional — required for queue reorder
    to_index       = d.get('to_index')        # optional — position within target queue
    if not all([from_s, from_b, to_s, to_b]):
        return jsonify({'error': 'missing fields'}), 400
    if from_b not in ('inprogress', 'queue') or to_b not in ('inprogress', 'queue'):
        return jsonify({'error': 'invalid bin type'}), 400

    try:
        # Hold the cache lock for the whole read-modify-write so this can't
        # interleave with the 60s poll thread (which also mutates these files).
        with _cache_lock:
            assignments = load_assignments()
            timers      = load_timers()
            prune_stale_assignments(assignments, timers)

            if from_s not in assignments:
                return jsonify({'error': f'unknown source station {from_s}'}), 400
            if to_s not in assignments:
                return jsonify({'error': f'unknown target station {to_s}'}), 400

            # Determine the PDO being moved
            if from_b == 'inprogress':
                pdo_id = assignments[from_s].get('inprogress')
            else:  # from queue
                pdo_id = pdo_id_hint
                if not pdo_id or pdo_id not in assignments[from_s].get('queue', []):
                    return jsonify({'error': 'source PDO not found in queue'}), 400

            if not pdo_id:
                return jsonify({'error': 'source bin empty'}), 400

            if to_b == 'inprogress' and assignments[to_s].get('inprogress'):
                return jsonify({'error': 'target in-progress occupied'}), 409

            # Remove from source
            if from_b == 'inprogress':
                assignments[from_s]['inprogress'] = None
                timers.pop(from_s, None)
            else:
                if pdo_id in assignments[from_s].get('queue', []):
                    assignments[from_s]['queue'].remove(pdo_id)

            # Insert at destination
            if to_b == 'inprogress':
                assignments[to_s]['inprogress'] = pdo_id
                timers[to_s] = {'started_at': datetime.now().isoformat(), 'fg_count': 0}
                workers_asgn = load_workers_assigned()
                if _count_active_workers(to_s, workers_asgn) == 0:
                    timers[to_s]['paused'] = True
            else:
                q = assignments[to_s].setdefault('queue', [])
                if isinstance(to_index, int) and 0 <= to_index <= len(q):
                    q.insert(to_index, pdo_id)
                else:
                    q.append(pdo_id)

            save_assignments(assignments)
            save_timers(timers)
            write_station_to_excel(pdo_id, to_s)
            update_split_station(pdo_id, to_s)
        return jsonify({'ok': True})
    except Exception as e:
        print(f'[api_move error] {e}')
        import traceback; traceback.print_exc()
        return jsonify({'error': str(e)}), 500

# ── API: reorder within a station's queue ─────────────────────────────────────
@app.route('/api/reorder_queue', methods=['POST'])
def api_reorder_queue():
    d = request.get_json(silent=True) or {}
    station, pdo_id, to_index = d.get('station'), d.get('pdo_id'), d.get('to_index')
    if not all([station, pdo_id]) or to_index is None:
        return jsonify({'error': 'missing fields'}), 400
    try:
        to_index = int(to_index)
    except (TypeError, ValueError):
        return jsonify({'error': 'invalid to_index'}), 400

    try:
        with _cache_lock:
            assignments = load_assignments()
            q = assignments.get(station, {}).get('queue', [])
            if pdo_id not in q:
                return jsonify({'error': 'pdo not in queue'}), 400

            q.remove(pdo_id)
            to_index = max(0, min(to_index, len(q)))
            q.insert(to_index, pdo_id)
            assignments[station]['queue'] = q

            save_assignments(assignments)
        return jsonify({'ok': True})
    except Exception as e:
        print(f'[api_reorder_queue error] {e}')
        return jsonify({'error': str(e)}), 500

# ── API: swap a queued PDO into in-progress, pushing the running one to front ───
@app.route('/api/replace_inprogress', methods=['POST'])
def api_replace_inprogress():
    """Promote a PDO (from this station's queue, another station, or the pool)
    into a station's in-progress slot, moving the CURRENT in-progress order to
    the FRONT of that station's queue (position 0). Backs the 'drag a queued
    order onto the running one to swap' gesture — the front-end confirms first."""
    d = request.get_json(silent=True) or {}
    station = d.get('station')
    pdo_id  = d.get('pdo_id')
    if not station or not pdo_id:
        return jsonify({'error': 'missing station or pdo_id'}), 400
    try:
        with _cache_lock:
            assignments = load_assignments()
            timers      = load_timers()
            prune_stale_assignments(assignments, timers)
            if station not in assignments:
                return jsonify({'error': f'unknown station {station}'}), 400

            old_ip = assignments[station].get('inprogress')
            if old_ip == pdo_id:
                return jsonify({'ok': True})  # already running here — nothing to do

            # Remove the incoming PDO from wherever it currently sits.
            for st, bins in assignments.items():
                if bins.get('inprogress') == pdo_id:
                    bins['inprogress'] = None
                    timers.pop(st, None)
                if pdo_id in bins.get('queue', []):
                    bins['queue'].remove(pdo_id)

            # Push the previously-running order to the FRONT of this station's queue.
            if old_ip:
                assignments[station].setdefault('queue', []).insert(0, old_ip)

            assignments[station]['inprogress'] = pdo_id
            workers_asgn = load_workers_assigned()
            timers[station] = {'started_at': datetime.now().isoformat(), 'fg_count': 0}
            if _count_active_workers(station, workers_asgn) == 0:
                timers[station]['paused'] = True

            save_assignments(assignments)
            save_timers(timers)
            write_station_to_excel(pdo_id, station)
            update_split_station(pdo_id, station)
        return jsonify({'ok': True})
    except Exception as e:
        print(f'[api_replace_inprogress error] {e}')
        import traceback; traceback.print_exc()
        return jsonify({'error': str(e)}), 500

# ── API: remove PDO from bin ──────────────────────────────────────────────────
@app.route('/api/remove', methods=['POST'])
def api_remove():
    d = request.get_json(silent=True) or {}
    station, bin_type = d.get('station'), d.get('bin')
    pdo_id_hint = d.get('pdo_id')  # required when removing from queue
    if not station or bin_type not in ('inprogress', 'queue'):
        return jsonify({'error': 'missing or invalid fields'}), 400

    try:
        with _cache_lock:
            assignments = load_assignments()
            timers      = load_timers()
            if station not in assignments:
                return jsonify({'error': f'unknown station {station}'}), 400

            advanced = None
            if bin_type == 'inprogress':
                pdo_id = assignments[station].get('inprogress')
                q = assignments[station].get('queue', [])
                if q:
                    # Auto-advance — same as completion/FG-driven advance, just triggered
                    # by a manual removal instead of the order actually finishing.
                    advanced = q.pop(0)
                    assignments[station]['inprogress'] = advanced
                    workers_asgn = load_workers_assigned()
                    worker_count = _count_active_workers(station, workers_asgn)
                    timers[station] = {'started_at': datetime.now().isoformat(), 'fg_count': 0}
                    if worker_count == 0:
                        timers[station]['paused'] = True
                else:
                    assignments[station]['inprogress'] = None
                    timers.pop(station, None)
            else:
                pdo_id = pdo_id_hint
                q = assignments[station].get('queue', [])
                if pdo_id in q:
                    q.remove(pdo_id)

            save_assignments(assignments)
            save_timers(timers)
            if pdo_id:
                write_station_to_excel(pdo_id, '')
                # Manual removal: suppress LP auto-requeue for this PDO until a
                # NEWER License Plate print is issued for it.
                _record_lp_suppression(pdo_id)
        return jsonify({'ok': True, 'advanced': advanced})
    except Exception as e:
        print(f'[api_remove error] {e}')
        import traceback; traceback.print_exc()
        return jsonify({'error': str(e)}), 500

# ── API: assign workers to station ───────────────────────────────────────────
@app.route('/api/workers/assign', methods=['POST'])
def api_workers_assign():
    d = request.get_json(silent=True) or {}
    station, worker_ids = d.get('station'), d.get('worker_ids', [])
    if not station:
        return jsonify({'error': 'missing station'}), 400
    if not isinstance(worker_ids, list):
        return jsonify({'error': 'worker_ids must be a list'}), 400
    try:
        with _cache_lock:
            workers_asgn = load_workers_assigned()
            # An operator can only be at ONE station at a time. Staffing them
            # here removes them from any other station they were on. This is the
            # authoritative guard against double-booking — e.g. dragging an
            # already-assigned worker out of the pool onto a second station
            # (a pool drag carries no source station, so the client couldn't
            # remove them from the first one).
            incoming = {str(w) for w in worker_ids}
            affected = set()
            for other_sid, ids in list(workers_asgn.items()):
                if other_sid == station:
                    continue
                kept = [w for w in ids if str(w) not in incoming]
                if len(kept) != len(ids):
                    workers_asgn[other_sid] = kept
                    affected.add(other_sid)
            workers_asgn[station] = worker_ids
            save_workers_assigned(workers_asgn)
            timers = load_timers()
            # Pause/resume off ACTIVE workers only — pre-assigning an upcoming
            # next-shift worker shouldn't start the clock before their shift.
            # Resync the target AND any station a worker was just pulled off of
            # (its active-worker count changed).
            sync_timer_pause_state(station, _count_active_workers(station, workers_asgn), timers)
            for sid in affected:
                sync_timer_pause_state(sid, _count_active_workers(sid, workers_asgn), timers)
            save_timers(timers)
        return jsonify({'ok': True})
    except Exception as e:
        print(f'[api_workers_assign error] {e}')
        return jsonify({'error': str(e)}), 500

# ── API: add/edit note on PDO ─────────────────────────────────────────────────
@app.route('/api/note', methods=['POST'])
def api_note():
    d = request.get_json(silent=True) or {}
    pdo_id, note = d.get('pdo_id'), d.get('note', '')
    if not pdo_id:
        return jsonify({'error': 'missing pdo_id'}), 400
    try:
        with _cache_lock:
            notes = load_json(f'notes_{AREA}.json', {})
            notes[pdo_id] = note if note is not None else ''
            save_json(f'notes_{AREA}.json', notes)
        return jsonify({'ok': True})
    except Exception as e:
        print(f'[api_note error] {e}')
        return jsonify({'error': str(e)}), 500

@app.route('/api/notes')
def api_notes():
    return jsonify(load_json(f'notes_{AREA}.json', {}))

# ── API: NG (defective) flag on a PDO ─────────────────────────────────────────
@app.route('/api/ng', methods=['GET', 'POST'])
def api_ng():
    """GET → {pdo_id: true} map of PDOs flagged NG (defective). POST {pdo_id, ng}
    sets/clears the flag. Persisted per area like notes."""
    if request.method == 'GET':
        return jsonify(load_json(f'ng_{AREA}.json', {}))
    d = request.get_json(silent=True) or {}
    pdo_id = d.get('pdo_id')
    if not pdo_id:
        return jsonify({'error': 'missing pdo_id'}), 400
    try:
        with _cache_lock:
            flags = load_json(f'ng_{AREA}.json', {})
            if d.get('ng'):
                flags[pdo_id] = True
            else:
                flags.pop(pdo_id, None)
            save_json(f'ng_{AREA}.json', flags)
        return jsonify({'ok': True})
    except Exception as e:
        print(f'[api_ng error] {e}')
        return jsonify({'error': str(e)}), 500

# ══════════════════════════════════════════════════════════════════════════════
# PRODUCTION TICKET / LABEL PRINTING  (BPT packaging ticket + Zebra box label)
# ══════════════════════════════════════════════════════════════════════════════
# Two print paths per PDO, replacing the Excel PrintLabel macro:
#   1. BPT packaging pick ticket  → warehouse Canon printer. Printed FROM THE
#      BROWSER: issuing returns a /bpt/<token> page that auto-opens the print
#      dialog, so the operator just confirms the Canon printer.
#   2. Box label + License Plate  → Zebra ZT421 (server-side RAW ZPL via win32).
# Issuing also QUEUES the PDO on the chosen station (this replaces the old
# LP-scan auto-queue as the placement signal). License Plate is a DUMMY for now.
import uuid as _uuid

_MAINDB_DEFAULT = r"\\npvshare\Data\06_Operation\02.Production\02.Component\01. General\MainDatabase\MainDatabase.xlsx"
_EMAIL_DEFAULT = [
    "Ginny.Do@northstar.vn", "Duc.nguyen@northstar.vn", "Nghia.nguyen@northstar.vn",
    "Tung.Le@Northstar.vn", "Richard.Pham@northstar.vn", "Tung.Hoang@northstar.vn",
    "Duong.Dong@northstar.vn", "Marcus.le@northstar.vn", "Jolie.Nguyen@northstar.vn",
    "Hai.le@northstar.vn", "Ryan.Nguyen@northstar.vn",
]

def _print_settings_defaults():
    return {
        'maindb_path': _MAINDB_DEFAULT,
        'zebra_printer': 'ZDesigner ZT421-300dpi ZPL',
        'canon_printer': r'\\10.191.56.14\vtn1prt-tanglung',
        'email_recipients': list(_EMAIL_DEFAULT),
        # Silent BPT paper path: render with headless Edge/Chrome, then push the
        # PDF to the named Canon queue with SumatraPDF. Blank paths = auto-detect.
        # If either piece is missing the board falls back to the browser tab.
        'silent_bpt': True,
        'sumatra_path': '',
        'browser_path': '',
        # Packaging materials charged per PALLET rather than per box, by code
        # prefix. Kept as settings so a new prefix is a config change, not a
        # code change (the business chose prefixes over a packaging_type table).
        'pallet_prefixes': list(_bpt_defaults()[0]),
        'crate_prefixes': list(_bpt_defaults()[1]),
    }

def _bpt_defaults():
    """(pallet_prefixes, crate_prefixes) from bpt.py, without importing it at
    module scope (bpt pulls in openpyxl/barcode, which the board shouldn't
    require just to boot)."""
    try:
        import bpt as _b
        return _b.PALLET_PREFIXES, _b.CRATE_PREFIXES
    except Exception:
        return ("PL", "SPL"), ("WC",)

def load_print_settings():
    s = _print_settings_defaults()
    s.update(load_json(f'print_settings_{AREA}.json', {}) or {})
    return s

# MainDatabase is big (~17k rows); cache the parsed index, refresh on a TTL or
# when the configured path changes.
_maindb_cache = {'db': None, 'path': None, 'ts': 0.0}
_MAINDB_TTL = 600  # 10 min

def get_maindb(force=False):
    import time as _t
    path = load_print_settings()['maindb_path']
    now = _t.time()
    c = _maindb_cache
    if (not force and c['db'] is not None and c['path'] == path
            and (now - c['ts']) < _MAINDB_TTL):
        return c['db']
    import bpt as _bpt
    db = _bpt.MainDB(path)      # raises if unreadable
    c.update(db=db, path=path, ts=now)
    return db

# Short-lived handoff cache: issue_ticket builds the BPT and stashes it here; the
# browser then GETs /bpt/<token> to render + print it.
_bpt_tickets = {}

def _issue_place_pdo(pdo_id, station):
    """Place a printed PDO on `station` (in-progress if free, else queue) — the
    print action's queue signal. Mirrors /api/assign placement."""
    assignments = load_assignments()
    timers = load_timers()
    if station not in assignments:
        assignments[station] = {'inprogress': None, 'queue': []}
    for st, bins in assignments.items():
        if bins.get('inprogress') == pdo_id:
            bins['inprogress'] = None
            timers.pop(st, None)
        if pdo_id in bins.get('queue', []):
            bins['queue'].remove(pdo_id)
    if not assignments[station]['inprogress']:
        assignments[station]['inprogress'] = pdo_id
        timers[station] = {'started_at': datetime.now().isoformat(), 'fg_count': 0}
        if _count_active_workers(station, load_workers_assigned()) == 0:
            timers[station]['paused'] = True
        placed = 'inprogress'
    else:
        assignments[station].setdefault('queue', []).append(pdo_id)
        placed = 'queue'
    save_assignments(assignments)
    save_timers(timers)
    try:
        write_station_to_excel(pdo_id, station)
        update_split_station(pdo_id, station)
    except Exception as e:
        print(f'[issue place] station write warn: {e}')
    return placed


@app.route('/api/printers')
def api_printers():
    import printing
    s = load_print_settings()
    ready, detail = printing.silent_ready(s.get('sumatra_path', ''), s.get('browser_path', ''))
    return jsonify({'printers': printing.list_printers(), 'win': printing.available(),
                    'silent_ready': ready, 'silent_detail': detail,
                    'browser': printing.find_browser(s.get('browser_path', '')),
                    'sumatra': printing.find_sumatra(s.get('sumatra_path', ''))})

@app.route('/api/print_settings', methods=['GET', 'POST'])
def api_print_settings():
    if request.method == 'GET':
        return jsonify(load_print_settings())
    d = request.get_json(silent=True) or {}
    s = load_print_settings()
    for k in ('maindb_path', 'zebra_printer', 'canon_printer', 'sumatra_path', 'browser_path'):
        if isinstance(d.get(k), str):
            s[k] = d[k].strip()
    if 'silent_bpt' in d:
        s['silent_bpt'] = bool(d['silent_bpt'])
    for k in ('pallet_prefixes', 'crate_prefixes'):
        if k in d:
            v = d[k]
            if isinstance(v, str):
                v = [x.strip().upper() for x in v.replace(';', ',').split(',') if x.strip()]
            if isinstance(v, list) and v:
                s[k] = v
    if 'email_recipients' in d:
        v = d['email_recipients']
        if isinstance(v, str):
            v = [x.strip() for x in v.replace('\n', ';').replace(',', ';').split(';') if x.strip()]
        if isinstance(v, list):
            s['email_recipients'] = v
    save_json(f'print_settings_{AREA}.json', s)
    _maindb_cache['db'] = None
    return jsonify({'ok': True, 'settings': s})

@app.route('/api/issue_ticket', methods=['POST'])
def api_issue_ticket():
    """Issue tickets/labels for a PDO and queue it on the station. Body:
    {pdo_id, station, pro_h, pro_m, do_bpt, do_boxlabel}."""
    d = request.get_json(silent=True) or {}
    pdo_id = d.get('pdo_id'); station = (d.get('station') or '').strip()
    if not pdo_id or not station:
        return jsonify({'error': 'missing pdo_id or station'}), 400
    do_bpt = d.get('do_bpt', True); do_box = d.get('do_boxlabel', True)
    ph = str(d.get('pro_h', '') or '').strip() or '00'
    pm = str(d.get('pro_m', '') or '').strip() or '00'
    pro_time = f'{ph.zfill(2)}:{pm.zfill(2)}'
    warnings = []
    try:
        import bpt as _bpt, boxlabel as _bl, printing as _pr
        with _cache_lock:
            p = pdo_by_id(pdo_id)
        if not p:
            return jsonify({'error': f'PDO {pdo_id} not on the active board'}), 404
        try:
            mdb = get_maindb()
        except Exception as e:
            return jsonify({'error': f'Cannot open MainDatabase: {e}'}), 500

        po = p['id']; part = p['part']; qty = p.get('qty', 0)
        dest = p.get('dest', ''); pii = p.get('pii_po', ''); vendor = p.get('vendor', '')
        prod_type = p.get('pack_type', ''); desc = p.get('desc', '')
        settings = load_print_settings()
        pal_pfx = settings.get('pallet_prefixes') or None
        crate_pfx = settings.get('crate_prefixes') or None

        # Resolve the packing standard ONCE and report it, so a wrong box count is
        # visible as "which qty/box did it use, from which market row" instead of
        # being silently baked into the ticket.
        bucket = _bpt.market_bucket(dest)
        std = _bpt.resolve_pack_std(mdb, part, bucket, pal_pfx)
        std_info = {'market': bucket, 'qty_per_box': std['qty_per_box'],
                    'box_per_pallet': std['box_per_pallet'], 'source': std['source'],
                    'buckets_available': std['buckets_available'], 'ok': std['ok'],
                    'error': std['error']}
        if not std['ok']:
            warnings.append(std['error'])
        elif std['source'] != 'exact':
            warnings.append(f"Packing standard taken from the {std['source']} row — "
                            f"no row for market '{bucket}'. Verify qty/box in MainDatabase.")

        lps = []
        boxes = []          # per-box outcome, shown in the confirmation dialog
        if do_box:
            r = _bl.build_box_labels(mdb, po=po, part=part, po_qty=qty, dest=dest,
                                     station=station, pii=pii, desc=desc,
                                     vendor_code=vendor, production_type=prod_type)
            if not r.get('ok'):
                warnings.append(f"Box label skipped: {r.get('error')}")
            else:
                for L in r['labels']:
                    lps.append(L['lp'])
                    ok, err = False, ''
                    if not _pr.available():
                        err = 'pywin32 not installed on this PC'
                    elif not settings.get('zebra_printer'):
                        err = 'no Zebra printer set in print settings'
                    else:
                        try:
                            _pr.print_zpl_raw(_bl.render_zpl(L['d']),
                                              settings['zebra_printer'], doc=f"LP {L['lp']}")
                            ok = True
                        except Exception as e:
                            err = str(e)
                    if not ok:
                        warnings.append(f"Zebra box {L['box_no']}: {err}")
                    boxes.append({'box_no': L['box_no'], 'qty': L['box_qty'],
                                  'lp': L['lp'], 'printed': ok, 'error': err})
                    # DUMMY LP → no [License Plate] INSERT yet (SQL phase).

        bpt_token = None
        bpt_pages = 0
        bpt_printed = False
        bpt_detail = ''
        if do_bpt:
            t = _bpt.build_bpt(mdb, fg=part, po=po, dest=dest, station=station,
                               pro_time=pro_time, po_qty=qty, fg_desc=desc,
                               pallet_prefixes=pal_pfx, crate_prefixes=crate_pfx)
            t['date'] = datetime.now().strftime('%d-%b-%y')
            t['issue_time'] = datetime.now().strftime('%H:%M')
            bpt_pages = t.get('total_boxes', 0)
            if not t.get('has_data'):
                warnings.append(
                    f"BPT has no rows to pick: {t.get('material_count', 0)} packaging BOM "
                    f"material(s) for market '{bucket}', qty/box {t.get('qty_per_box_std', 0)}")
            bpt_token = _uuid.uuid4().hex
            _bpt_tickets[bpt_token] = t
            if len(_bpt_tickets) > 50:
                for k in list(_bpt_tickets)[:-50]:
                    _bpt_tickets.pop(k, None)
            # Silent paper path (headless render → named Canon queue). Falls back
            # to the browser tab, which is what bpt_printed=False tells the client.
            if settings.get('silent_bpt') and bpt_pages:
                html = _bpt.render_html(t)
                bpt_printed, bpt_detail = _pr.print_html_silent(
                    html, settings.get('canon_printer', ''),
                    settings.get('sumatra_path', ''), settings.get('browser_path', ''))
                if not bpt_printed:
                    warnings.append(f"Silent BPT print unavailable ({bpt_detail}) — "
                                    f"opening the print dialog instead")

        with _cache_lock:
            placed = _issue_place_pdo(pdo_id, station)

        return jsonify({'ok': True, 'lps': lps, 'boxes': boxes, 'std': std_info,
                        'bpt_token': bpt_token, 'bpt_pages': bpt_pages,
                        'bpt_printed': bpt_printed, 'bpt_detail': bpt_detail,
                        'do_bpt': bool(do_bpt), 'do_box': bool(do_box),
                        'pdo_id': pdo_id, 'part': part, 'po_qty': qty,
                        'placed': placed, 'warnings': warnings})
    except Exception as e:
        print(f'[api_issue_ticket error] {e}')
        return jsonify({'error': str(e)}), 500

@app.route('/bpt/<token>')
def bpt_page(token):
    import bpt as _bpt
    t = _bpt_tickets.get(token)
    if not t:
        return "Ticket expired — re-issue it from the board.", 404
    html = _bpt.render_html(t)
    return html.replace('</body>', '<script>window.addEventListener("load",()=>setTimeout(()=>window.print(),300))</script></body>')


# ── Startup config window (Tkinter) ─────────────────────────────────────────────
# Shown only when running `python app.py` directly (not on import/test-client use).
# Lets the user verify/change the plan Excel + attendance Excel paths, pings each
# for existence + readability, and only proceeds to start Flask once confirmed.
STARTUP_CONFIG_PATH = os.path.join(BASE_DIR, 'startup_config.json')

def _load_startup_config():
    if os.path.exists(STARTUP_CONFIG_PATH):
        try:
            with open(STARTUP_CONFIG_PATH, encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass
    return {}

def _validate_excel_path(path, required_sheet=None):
    """Returns (ok: bool, message: str). Reads the file via the SAME
    copy-then-open path the running server uses (_safe_load_workbook), NOT a
    direct openpyxl open. That matters: a file merely open in Excel isn't
    exclusively locked (Excel only holds the lock during the brief save), so the
    server can read it via a temp copy — and so can this check. Validating the
    same way means Start is no longer falsely blocked just because someone has
    the workbook open; genuinely missing/corrupted files still fail."""
    if not path:
        return False, '✗ No path set'
    if not os.path.exists(path):
        return False, '✗ File not found at this path'
    tmp_path = None
    try:
        wb, tmp_path = _safe_load_workbook(path, read_only=True, data_only=True)
        sheets = wb.sheetnames
        wb.close()
        if required_sheet and required_sheet not in sheets:
            return False, f'⚠ Opened OK, but sheet "{required_sheet}" not found ({len(sheets)} sheet(s) present)'
        note = ' (open in Excel — read via copy)' if _is_excel_locked(path) else ''
        return True, f'✓ Found and readable — {len(sheets)} sheet(s){note}'
    except zipfile.BadZipFile:
        return False, '⚠ Found, but could not open — the file appears corrupted'
    except Exception as e:
        return False, f'⚠ Found, but error opening: {e}'
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try: os.remove(tmp_path)
            except: pass

def run_startup_and_console(start_backend):
    """Single Tk root for the whole interactive session. Shows the setup
    screen first, then — WITHOUT ever destroying/recreating the Tk root —
    tears down the setup widgets and builds the log console into the same
    root once Start is confirmed, then calls start_backend().

    Creating a second tk.Tk() after destroying the first is a known Tkinter
    footgun: leftover Variable objects from the dead interpreter get
    garbage-collected while the new one's mainloop is running, throwing
    'main thread is not in main loop' repeatedly. One root for the process's
    whole life avoids it entirely.

    Any failure to load Tkinter itself (missing on this system, no display,
    etc.) is caught and treated as 'proceed with existing/default paths' —
    a missing GUI toolkit should never block the server from starting.
    Returns True if the server was started, False if cancelled before that."""
    global EXCEL_PLAN, AREA
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox
    except Exception as e:
        print(f'[startup config] Tkinter unavailable ({e}) — skipping setup window, using existing configured paths.')
        start_backend()
        return True

    saved = _load_startup_config()
    initial_plan = saved.get('excel_plan', EXCEL_PLAN)
    initial_att  = saved.get('attendance_pk', CAPACITY_FILES.get(AREA, ''))

    result = {'started': False}

    try:
        root = tk.Tk()
        root.title('Production Control Board — Startup Setup')
        root.geometry('940x640')
        root.minsize(760, 520)
        root.resizable(True, True)

        bg, fg, accent, entry_bg, muted = '#171a24', '#f7f8fc', '#2fdb96', '#2a2f3d', '#b0b6c9'
        root.configure(bg=bg)

        setup_frame = tk.Frame(root, bg=bg)
        setup_frame.pack(fill='both', expand=True)

        tk.Label(setup_frame, text='Production Control Board — Startup Setup', bg=bg, fg=fg,
                 font=('Segoe UI', 17, 'bold')).pack(pady=(26, 6))
        tk.Label(setup_frame, text='Verify the file paths below, then start the server.',
                 bg=bg, fg=muted, font=('Segoe UI', 11)).pack(pady=(0, 2))
        tk.Label(setup_frame, text=f'Build {BUILD_VERSION}  ·  {BUILD_NOTE}',
                 bg=bg, fg=accent, font=('Consolas', 9)).pack(pady=(0, 18))

        plan_var = tk.StringVar(value=initial_plan)
        att_var  = tk.StringVar(value=initial_att)
        plan_status = tk.StringVar(value='')
        att_status  = tk.StringVar(value='')
        status_labels = {}

        def make_row(label_text, var, status_var, key):
            frame = tk.Frame(setup_frame, bg=bg)
            frame.pack(fill='x', padx=36, pady=12)
            tk.Label(frame, text=label_text, bg=bg, fg=fg, font=('Segoe UI', 12, 'bold'), anchor='w').pack(fill='x')
            row = tk.Frame(frame, bg=bg)
            row.pack(fill='x', pady=(6, 0))
            entry = tk.Entry(row, textvariable=var, bg=entry_bg, fg=fg, insertbackground=fg,
                              relief='flat', font=('Consolas', 11))
            entry.pack(side='left', fill='x', expand=True, ipady=8)

            def browse():
                path = filedialog.askopenfilename(
                    filetypes=[('Excel macro files', '*.xlsm'), ('Excel files', '*.xlsx'), ('All files', '*.*')])
                if path:
                    var.set(path)
                    do_validate()

            tk.Button(row, text='Browse…', command=browse, bg=entry_bg, fg=fg,
                      relief='flat', font=('Segoe UI', 11), padx=16, pady=6, cursor='hand2').pack(side='left', padx=(10, 0))

            status_lbl = tk.Label(frame, textvariable=status_var, bg=bg, fg=muted,
                                   font=('Segoe UI', 10), anchor='w')
            status_lbl.pack(fill='x', pady=(7, 0))
            status_labels[key] = status_lbl

        make_row('Production plan Excel (.xlsm)', plan_var, plan_status, 'plan')
        make_row('Attendance matrix Excel (.xlsm)', att_var, att_status, 'att')

        # ── CS 1.6-style blocky loading bar ─────────────────────────────────
        # Purely cosmetic homage, but also functionally useful: opening these
        # files is a real network/disk hit (UNC paths, sometimes-slow shares),
        # so a marching segmented bar makes clear the window hasn't frozen
        # instead of validate()/Start Server just hanging silently.
        N_SEGMENTS = 28
        BAR_W, BAR_H = 700, 26
        progress_wrap = tk.Frame(setup_frame, bg=bg)
        progress_canvas = tk.Canvas(progress_wrap, width=BAR_W, height=BAR_H, bg='#0c0e14',
                                     highlightthickness=1, highlightbackground=muted)
        progress_canvas.pack()
        progress_label = tk.StringVar(value='')
        tk.Label(progress_wrap, textvariable=progress_label, bg=bg, fg=accent,
                 font=('Consolas', 10, 'bold')).pack(pady=(6, 0))

        def draw_bar(fraction):
            progress_canvas.delete('all')
            gap = 3
            seg_w = (BAR_W - gap * (N_SEGMENTS - 1)) / N_SEGMENTS
            filled = int(round(max(0, min(1, fraction)) * N_SEGMENTS))
            for i in range(N_SEGMENTS):
                x0 = i * (seg_w + gap)
                x1 = x0 + seg_w
                color = accent if i < filled else '#242938'
                progress_canvas.create_rectangle(x0, 2, x1, BAR_H - 2, fill=color, outline='')

        _VALIDATION_STAGES = [
            ('Connecting to plan file location…',        0.12),
            ('Reading production plan workbook…',         0.35),
            ('Checking sheet & area…',                    0.52),
            ('Connecting to attendance file location…',   0.68),
            ('Reading attendance workbook…',               0.90),
        ]

        def run_validation_animated(on_done):
            """Runs the real _validate_excel_path() calls on a background
            thread while the bar above marches through the stages in the
            foreground — real result is applied whenever the thread actually
            finishes, the animation just fills the wait."""
            result_v = {}

            def worker():
                ok1, msg1 = _validate_excel_path(plan_var.get(), required_sheet=PLAN_SHEET)
                filename_ok = True
                if ok1:
                    area, area_err = detect_area_filter(plan_var.get())
                    if area is None:
                        filename_ok = False
                        ok1 = False
                        msg1 = f'🚫 {area_err}'
                ok2, msg2 = _validate_excel_path(att_var.get())
                result_v.update(ok1=ok1, msg1=msg1, ok2=ok2, msg2=msg2, filename_ok=filename_ok, done=True)

            threading.Thread(target=worker, daemon=True).start()
            progress_wrap.pack(pady=(6, 4))
            stage_idx = 0

            def animate():
                nonlocal stage_idx
                if result_v.get('done'):
                    draw_bar(1.0)
                    progress_label.set('✓ Done')
                    root.after(220, lambda: finish())
                    return
                text, frac = _VALIDATION_STAGES[stage_idx]
                progress_label.set(text)
                draw_bar(frac)
                if stage_idx < len(_VALIDATION_STAGES) - 1:
                    stage_idx += 1
                root.after(150, animate)

            def finish():
                progress_wrap.pack_forget()
                on_done(result_v['ok1'], result_v['msg1'], result_v['ok2'], result_v['msg2'], result_v['filename_ok'])

            animate()

        def apply_validation_results(ok1, msg1, ok2, msg2):
            plan_status.set(msg1)
            status_labels['plan'].configure(fg=accent if ok1 else '#ff6363')
            att_status.set(msg2)
            status_labels['att'].configure(fg=accent if ok2 else '#ff6363')

        def set_buttons_enabled(enabled):
            state = 'normal' if enabled else 'disabled'
            validate_btn.configure(state=state)
            start_btn.configure(state=state)
            cancel_btn.configure(state=state)

        def do_validate(then=None):
            set_buttons_enabled(False)
            def on_done(ok1, msg1, ok2, msg2, filename_ok):
                apply_validation_results(ok1, msg1, ok2, msg2)
                set_buttons_enabled(True)
                if then:
                    then(ok1, ok2, filename_ok)
            run_validation_animated(on_done)

        validate_btn = tk.Button(setup_frame, text='🔄  Test / validate paths', command=lambda: do_validate(),
                                  bg=entry_bg, fg=fg, relief='flat', font=('Segoe UI', 11), padx=18, pady=8, cursor='hand2')
        validate_btn.pack(pady=(14, 6))

        # Dev/test mode toggle — when on, no big-screen feed and SQL writes are dry-run.
        dev_var = tk.BooleanVar(value=DEV_MODE)
        dev_chk = tk.Checkbutton(
            setup_frame, variable=dev_var,
            text='🧪  Dev / test mode  —  don\'t feed the big screen, and make all SQL writes dry-run',
            bg=bg, fg='#ffcf6b', selectcolor=entry_bg, activebackground=bg, activeforeground='#ffcf6b',
            font=('Segoe UI', 10, 'bold'), anchor='w', cursor='hand2')
        dev_chk.pack(pady=(16, 0))

        btn_row = tk.Frame(setup_frame, bg=bg)
        btn_row.pack(pady=(10, 0))

        def start_console_phase():
            """Transition in place: tear down the setup widgets and build the
            log console into the SAME root/mainloop — no new Tk() instance,
            no nested mainloop() call, so nothing from the setup screen
            outlives its interpreter incorrectly."""
            setup_frame.destroy()
            _build_console_ui(root, start_backend)

        def on_start():
            def after_validate(ok1, ok2, filename_ok):
                global EXCEL_PLAN, AREA
                if not filename_ok:
                    # Hard block — no "continue anyway" bypass for this one, since
                    # showing the wrong area's orders is worse than not starting.
                    _, area_err = detect_area_filter(plan_var.get())
                    messagebox.showerror(
                        'Invalid plan file name',
                        f'{area_err}\n\n'
                        'The server cannot start until this is fixed — rename the file '
                        'to start with Packaging, AssyD, or RawPart, or point to the correct file.'
                    )
                    return
                if not (ok1 and ok2):
                    # Hard block, same as the filename check — no "continue
                    # anyway" bypass. A red status means the file couldn't be
                    # reached/read at all, so starting anyway would just mean
                    # a board with silently wrong/stale data from minute one.
                    bad = []
                    if not ok1: bad.append(f'• Production plan Excel — {plan_status.get()}')
                    if not ok2: bad.append(f'• Attendance matrix Excel — {att_status.get()}')
                    messagebox.showerror(
                        'Cannot start — file check failed',
                        'One or more files could not be validated:\n\n'
                        + '\n'.join(bad) +
                        '\n\nFix the path (or the file/network issue) and click '
                        '"Test / validate paths" again before starting.'
                    )
                    return
                try:
                    with open(STARTUP_CONFIG_PATH, 'w', encoding='utf-8') as f:
                        json.dump({'excel_plan': plan_var.get(), 'attendance_pk': att_var.get()}, f, ensure_ascii=False, indent=2)
                except Exception as e:
                    print(f'[startup config] could not save config: {e}')
                EXCEL_PLAN = plan_var.get()
                AREA = compute_area_code()  # recompute — the GUI may have pointed at a different area's file
                CAPACITY_FILES[AREA] = att_var.get()
                global DEV_MODE
                DEV_MODE = bool(dev_var.get())
                result['started'] = True
                start_console_phase()
            do_validate(then=after_validate)

        def on_cancel():
            root.destroy()

        cancel_btn = tk.Button(btn_row, text='Cancel', command=on_cancel, bg=entry_bg, fg=fg,
                                relief='flat', font=('Segoe UI', 12), padx=28, pady=10, cursor='hand2')
        cancel_btn.pack(side='left', padx=8)
        start_btn = tk.Button(btn_row, text='▶  Start Server', command=on_start, bg=accent, fg='#0a1a12',
                               relief='flat', font=('Segoe UI', 12, 'bold'), padx=28, pady=10, cursor='hand2')
        start_btn.pack(side='left', padx=8)

        do_validate()  # auto-check on open
        root.mainloop()
    except Exception as e:
        print(f'[startup config] GUI error ({e}) — proceeding with existing configured paths.')
        return True

    return result['started']

# ── In-app debug console (Tkinter) ──────────────────────────────────────────────
# Replaces the bare cmd.exe window entirely. From the moment this window opens,
# stdout/stderr (every print() in this file, Flask/werkzeug's own request
# logging, background thread output) is queued and drained into a styled Text
# widget instead of a terminal — so launching via pythonw.exe (no console at
# all) still gives a live, readable log. Flask itself moves onto a background
# thread so Tkinter's mainloop (which must own the main thread) can run here.
class _QueueWriter:
    """A stdout/stderr replacement. Only ever pushes onto a thread-safe queue —
    safe to call from any thread (Flask's request thread, poll loops, etc.).
    Optionally mirrors to the real original stream too, so running via a
    normal python.exe console still shows output there as well."""
    def __init__(self, q, mirror=None):
        self._q = q
        self._mirror = mirror

    def write(self, text):
        # Werkzeug's dev server (and other libs) occasionally write raw
        # bytes rather than str — normalize here so nothing downstream
        # (the queue, append_console's .lower()) ever has to guess.
        if isinstance(text, bytes):
            text = text.decode('utf-8', errors='replace')
        if text:
            self._q.put(text)
        if self._mirror:
            try:
                self._mirror.write(text)
            except Exception:
                pass

    def flush(self):
        if self._mirror:
            try:
                self._mirror.flush()
            except Exception:
                pass

def _build_console_ui(root, start_backend):
    """Builds the log console UI into an ALREADY-EXISTING Tk root (the same
    one the setup screen just vacated) and kicks off start_backend(). Never
    creates a new Tk() instance and never calls mainloop() itself — the
    caller's root.mainloop() (already running, since this executes inside a
    button callback) keeps driving the event loop from here on."""
    import tkinter as tk
    from tkinter import messagebox

    root.title(f'Production Control Board — {AREA} — Server Console')
    root.geometry('980x580')
    root.minsize(640, 360)

    log_queue = queue.Queue()
    sys.stdout = _QueueWriter(log_queue, mirror=sys.__stdout__)
    sys.stderr = _QueueWriter(log_queue, mirror=sys.__stderr__)

    bg, fg, accent, entry_bg, muted, err_c, warn_c = \
        '#0c0e14', '#e8eaf2', '#2fdb96', '#171a24', '#7d8399', '#ff6363', '#e8b84b'
    root.configure(bg=bg)

    header = tk.Frame(root, bg=entry_bg)
    header.pack(fill='x')
    tk.Label(header, text=f'  Production Control Board — {AREA}', bg=entry_bg, fg=fg,
             font=('Segoe UI', 12, 'bold'), anchor='w').pack(side='left', pady=10)
    status_var = tk.StringVar(value='●  LOADING…')
    status_lbl = tk.Label(header, textvariable=status_var, bg=entry_bg, fg=warn_c,
                           font=('Consolas', 11, 'bold'), anchor='e')
    status_lbl.pack(side='right', padx=16)

    toolbar = tk.Frame(root, bg=bg)
    toolbar.pack(fill='x', padx=10, pady=(8, 0))
    autoscroll_var = tk.BooleanVar(value=True)
    tk.Checkbutton(toolbar, text='Auto-scroll', variable=autoscroll_var, bg=bg, fg=muted,
                   selectcolor=entry_bg, activebackground=bg, activeforeground=fg,
                   font=('Segoe UI', 9)).pack(side='left')

    def open_dashboard():
        try:
            webbrowser.open('http://127.0.0.1:5050/', new=2)
        except Exception as e:
            messagebox.showerror('Could not open browser', str(e))

    open_dash_btn = tk.Button(toolbar, text='🌐  Open Dashboard', command=open_dashboard, bg=entry_bg, fg=muted,
                               relief='flat', font=('Segoe UI', 9, 'bold'), padx=12, pady=3,
                               cursor='hand2', state='disabled')
    open_dash_btn.pack(side='left', padx=(10, 0))

    def mark_ready():
        # Called via root.after() from the backend thread's on_ready callback
        # once initialize_app() has actually finished (Excel/SQL init done) —
        # NOT just as soon as the background thread was launched, so
        # 'RUNNING' actually means running, not merely started.
        status_var.set('●  RUNNING')
        status_lbl.configure(fg=accent)
        open_dash_btn.configure(fg=accent, state='normal')

    def shutdown_server():
        if messagebox.askyesno('Shut down server?',
                                'This will stop the Production Control Board server and close this window. Continue?'):
            os._exit(0)

    tk.Button(toolbar, text='⏻  Shutdown', command=shutdown_server, bg=entry_bg, fg=err_c,
              relief='flat', font=('Segoe UI', 9, 'bold'), padx=12, pady=3, cursor='hand2').pack(side='right')

    def clear_console():
        console_text.configure(state='normal')
        console_text.delete('1.0', 'end')
        console_text.configure(state='disabled')

    tk.Button(toolbar, text='Clear', command=clear_console, bg=entry_bg, fg=fg,
              relief='flat', font=('Segoe UI', 9), padx=12, pady=3, cursor='hand2').pack(side='right', padx=(0, 8))

    text_frame = tk.Frame(root, bg=bg)
    text_frame.pack(fill='both', expand=True, padx=10, pady=10)
    scrollbar = tk.Scrollbar(text_frame)
    scrollbar.pack(side='right', fill='y')
    console_text = tk.Text(text_frame, bg=bg, fg=fg, insertbackground=fg,
                            font=('Consolas', 10), wrap='word', relief='flat',
                            yscrollcommand=scrollbar.set, state='disabled', padx=8, pady=6)
    console_text.pack(side='left', fill='both', expand=True)
    scrollbar.config(command=console_text.yview)
    console_text.tag_configure('err', foreground=err_c)
    console_text.tag_configure('warn', foreground=warn_c)
    console_text.tag_configure('ok', foreground=accent)

    def append_console(text):
        if not isinstance(text, str):
            text = str(text)
        low = text.lower()
        tag = None
        if 'refusing' in low or '✗' in text or '🚫' in text or 'error' in low or 'traceback' in low:
            tag = 'err'
        elif 'warning' in low or '⚠' in text:
            tag = 'warn'
        elif text.startswith('[startup]') or '✓' in text:
            tag = 'ok'
        console_text.configure(state='normal')
        console_text.insert('end', text, tag)
        if autoscroll_var.get():
            console_text.see('end')
        console_text.configure(state='disabled')

    def drain_queue():
        # A single malformed line must never be able to stall the whole
        # console — without this, one bad append_console() call raises out
        # of the loop, skips the root.after() reschedule below, and the
        # console silently freezes forever while the real server keeps
        # running (as happened with a stray bytes object from Werkzeug).
        try:
            while True:
                try:
                    append_console(log_queue.get_nowait())
                except queue.Empty:
                    break
                except Exception as e:
                    try:
                        append_console(f'[console] (dropped a malformed log line: {e})\n')
                    except Exception:
                        pass
        finally:
            root.after(80, drain_queue)

    def on_close():
        if messagebox.askyesno('Close console?', 'Closing this window will stop the server. Continue?'):
            os._exit(0)
    root.protocol('WM_DELETE_WINDOW', on_close)

    drain_queue()
    root.after(150, lambda: start_backend(on_ready=lambda: root.after(0, mark_ready)))

# ── Startup ────────────────────────────────────────────────────────────────────
# Apply any previously-saved path config (from an earlier run of the setup
# window) so even a headless import/test-client run reflects the last-chosen
# paths — the GUI below is where the user changes them, this just persists it.
_saved_cfg = _load_startup_config()
if _saved_cfg.get('excel_plan'):
    EXCEL_PLAN = _saved_cfg['excel_plan']
if _saved_cfg.get('attendance_pk'):
    # NOTE: 'attendance_pk' is just the config file's internal key name (kept
    # for backward compatibility with already-saved configs) — it holds
    # whichever area's attendance file was configured, not necessarily
    # Packaging's. Applied to CAPACITY_FILES[AREA] below once AREA is known.
    _saved_attendance_override = _saved_cfg['attendance_pk']
else:
    _saved_attendance_override = None

# AREA must be computed AFTER EXCEL_PLAN reflects any saved override above —
# everything else (data file names, attendance lookup, station defaults)
# depends on this being correct for whichever area's Excel this instance
# was actually pointed at.
AREA = compute_area_code()
if _saved_attendance_override:
    CAPACITY_FILES[AREA] = _saved_attendance_override

def initialize_app():
    """The heavy init pass — Excel reads, SQL queries, background poll thread.
    Runs exactly once: either right after the setup window confirms (so it
    reflects whatever paths the user just picked), or immediately on import
    for non-interactive use (test clients, WSGI servers) where there's no
    GUI to wait for."""
    print(f'[startup] ===== Production Control Board  build {BUILD_VERSION}  ({BUILD_NOTE}) =====')
    if DEV_MODE:
        print('[startup] *** DEV/TEST MODE ON — big-screen feed DISABLED, SQL writes are DRY-RUN (nothing written to production) ***')
    print(f'[startup] BASE_DIR: {BASE_DIR}')
    print(f'[startup] DATA_DIR: {DATA_DIR}')
    print(f'[startup] EXCEL_PLAN exists: {os.path.exists(EXCEL_PLAN)} → {EXCEL_PLAN}')
    print(f'[startup] DEV_PLAN   exists: {os.path.exists(DEV_PLAN)} → {DEV_PLAN}')
    print(f'[startup] Attendance {AREA} file exists: {os.path.exists(CAPACITY_FILES.get(AREA,""))} → {CAPACITY_FILES.get(AREA)}')
    print(f'[startup] pyodbc available: {PYODBC_AVAILABLE}')
    refresh_cache(force=True)
    refresh_fg_and_advance()  # initial FG pull right away, not just on the 60s cycle
    print(f'[startup] PDOs loaded: {len(_pdo_cache)} (raw from Excel: {len(_raw_pdo_cache)})')
    print(f'[startup] Workers currently on-shift: {len(_worker_cache)}')
    print(f'[startup] Stations loaded: {_station_cache}')
    print(f'[startup] TimeStudy status: {_last_ts_status} | FG status: {_last_fg_status}')
    threading.Thread(target=_fg_poll_loop, daemon=True).start()
    threading.Thread(target=_dashboard_write_loop, daemon=True).start()
    print(f'[startup] Shop floor dashboard feed: writing to {DASHBOARD_STATUS_DIR} every {DASHBOARD_WRITE_INTERVAL_SEC}s')

if __name__ == '__main__':
    # Get rid of the cmd.exe window right away — before the setup screen
    # even opens — so it never flashes visibly at any point. Relaunching via
    # pythonw.exe is the reliable path; ShowWindow-hiding only runs as a
    # fallback if that's not possible on this machine.
    _relaunch_without_console()
    _hide_console_window()

    # Block a second copy of this install BEFORE anything heavy (or any file
    # write / poll thread) starts, so two instances can never corrupt the shared
    # data files or fight over port 5050.
    if not acquire_single_instance_lock():
        _notify_already_running()
        sys.exit(0)

    def _start_backend(on_ready=None):
        # Runs on a background thread. Final safety net — even if the GUI
        # was skipped entirely (e.g. Tkinter unavailable on this machine) or
        # somehow bypassed, the server must never start with a plan file
        # whose name doesn't tell us which area it's for.
        _area, _area_err = detect_area_filter(EXCEL_PLAN)
        if _area is None:
            print(f'[startup] REFUSING TO START: {_area_err}')
            print('[startup] Fix EXCEL_PLAN (must start with Packaging / AssyD / RawPart) and try again.')
            return
        initialize_app()  # the actual heavy lifting — Excel/SQL/attendance reads
        if on_ready:
            on_ready()  # fires only once init has genuinely finished, not when the thread merely started
        # debug=False in production: the Werkzeug interactive debugger is a
        # remote-code-execution surface if the port is ever reachable, and it's
        # of no use here since logs already stream to the in-app console.
        # threaded=True keeps the board responsive when several screens poll at
        # once; use_reloader=False is required (Flask runs on a background
        # thread, and the reloader can only arm signals from the main thread).
        app.run(debug=False, port=5050, use_reloader=False, threaded=True)

    def _launch_backend_thread(on_ready=None):
        threading.Thread(target=_start_backend, kwargs={'on_ready': on_ready}, daemon=True).start()

    # Interactive run: show the setup window FIRST, before any Excel/SQL reads —
    # previously this init ran unconditionally at import time, so the console
    # would churn through a full (wasted) init pass using the old default paths
    # before the window even appeared. Now nothing heavy happens until Start
    # is clicked, and it only runs once, using whatever paths were confirmed.
    # One Tk root drives both the setup screen and (once confirmed) the log
    # console — see run_startup_and_console() for why that matters.
    proceed = run_startup_and_console(_launch_backend_thread)
    if not proceed:
        print('[startup] Cancelled from setup window — server not started.')
else:
    # Imported (test client, WSGI, etc.) — no GUI to wait for, init immediately.
    initialize_app()
